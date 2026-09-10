from pathlib import Path
import hashlib
import shutil
import tempfile
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REFERENCE = Path(r"E:\2026\Cải tiến\VWork\VWork_feedback-mockup_code-UC-van-hanh-dao-tao\VWork_feedback-mockup_code-UC-van-hanh-dao-tao_06_can-sua.docx")
SOURCE_IMAGE = Path(r"C:\Users\admin\AppData\Local\Temp\codex-clipboard-455b53ae-77af-412e-92e8-5f2e96aa6527.png")
OUTPUT = Path(r"E:\2026\Cải tiến\VWork\VWork_feedback-mockup_code-UC-van-hanh-dao-tao\VWork_feedback-mockup_code-UC-van-hanh-dao-tao_07_can-xac-nhan.docx")
REFERENCE_SHA256 = "7EAC30FF31005599370DCC903831418943054FEF23B5EDE0E5B94DB76D98DA68"


def set_font(run, name="Arial", size=12, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def clear_body_keep_section(doc):
    body = doc._element.body
    section = body.sectPr
    for child in list(body):
        if child is not section:
            body.remove(child)


def ensure_style(doc, name, style_type, size, bold=False):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, style_type)
    style.font.name = "Arial"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    p_pr = style._element.get_or_add_pPr()
    for border in p_pr.findall(qn("w:pBdr")):
        p_pr.remove(border)
    return style


def apply_crop(inline_shape, left, top, right, bottom):
    # Values are percentages of the original image expressed in 1/1000 percent.
    blip_fill = inline_shape._inline.graphic.graphicData.pic.blipFill
    src_rect = blip_fill.find(qn("a:srcRect"))
    if src_rect is None:
        src_rect = OxmlElement("a:srcRect")
        blip_fill.insert(1, src_rect)
    src_rect.set("l", str(left))
    src_rect.set("t", str(top))
    src_rect.set("r", str(right))
    src_rect.set("b", str(bottom))


def remove_unreferenced_media(docx_path):
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    with zipfile.ZipFile(docx_path, "r") as source:
        parts = {name: source.read(name) for name in source.namelist()}

    document = ET.fromstring(parts["word/document.xml"])
    used_ids = {
        node.attrib[f"{{{office_rel_ns}}}embed"]
        for node in document.findall(f".//{{{drawing_ns}}}blip")
        if f"{{{office_rel_ns}}}embed" in node.attrib
    }

    rels_name = "word/_rels/document.xml.rels"
    rels = ET.fromstring(parts[rels_name])
    for rel in list(rels):
        if rel.attrib.get("Type", "").endswith("/image") and rel.attrib.get("Id") not in used_ids:
            rels.remove(rel)
    parts[rels_name] = ET.tostring(rels, encoding="utf-8", xml_declaration=True)

    referenced_media = set()
    for name, data in parts.items():
        if not name.endswith(".rels"):
            continue
        try:
            tree = ET.fromstring(data)
        except ET.ParseError:
            continue
        rel_dir = Path(name).parent.parent
        for rel in tree.findall(f"{{{rel_ns}}}Relationship"):
            if not rel.attrib.get("Type", "").endswith("/image"):
                continue
            target = rel.attrib.get("Target", "")
            if target.startswith("../"):
                resolved = (rel_dir / target).as_posix()
                while "/../" in resolved:
                    before, after = resolved.split("/../", 1)
                    resolved = before.rsplit("/", 1)[0] + "/" + after
            else:
                resolved = (rel_dir / target).as_posix()
            referenced_media.add(resolved)

    for name in list(parts):
        if name.startswith("word/media/") and name not in referenced_media:
            del parts[name]

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=docx_path.parent) as handle:
        temp_path = Path(handle.name)
    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name, data in parts.items():
                target.writestr(name, data)
        temp_path.replace(docx_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def main():
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    if not SOURCE_IMAGE.exists():
        raise FileNotFoundError(SOURCE_IMAGE)
    actual_hash = hashlib.sha256(REFERENCE.read_bytes()).hexdigest().upper()
    if actual_hash != REFERENCE_SHA256:
        raise RuntimeError(f"Reference hash mismatch: {actual_hash}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    title_style = ensure_style(doc, "Title", WD_STYLE_TYPE.PARAGRAPH, 20, True)
    title_style.paragraph_format.space_after = Pt(10)
    title_style.paragraph_format.keep_with_next = True
    heading_style = ensure_style(doc, "Heading 1", WD_STYLE_TYPE.PARAGRAPH, 14, True)
    heading_style.paragraph_format.space_before = Pt(12)
    heading_style.paragraph_format.space_after = Pt(6)
    heading_style.paragraph_format.keep_with_next = True

    title = doc.add_paragraph(style="Title")
    title.add_run("Feedback cần sửa VWork")

    intro = doc.add_paragraph()
    intro.add_run(
        "Tài liệu ghi nhận feedback về luồng gán vai trò trong phần Ekip khóa học."
    )

    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run("FB 01 Cải thiện khả năng tìm thấy thao tác Gán vai trò")

    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.add_run("Vị trí hiện tại của thao tác Gán vai trò khó tìm. Cần điều chỉnh để người dùng dễ nhận biết và thực hiện thao tác này trong luồng quản lý ekip khóa học.")

    evidence = doc.add_paragraph()
    evidence.alignment = WD_ALIGN_PARAGRAPH.CENTER
    evidence.paragraph_format.space_before = Pt(4)
    evidence.paragraph_format.keep_together = True
    shape = evidence.add_run().add_picture(str(SOURCE_IMAGE), width=Inches(6.1))
    shape._inline.docPr.set("title", "Ảnh minh họa vị trí Gán vai trò")
    shape._inline.docPr.set(
        "descr",
        "Ảnh trao đổi và giao diện Ekip khóa học cho thấy nút Gán vai trò nằm cạnh ô chọn tài khoản và vai trò.",
    )
    # Focus on the visible feedback message and its attached VWork screens.
    apply_crop(shape, 23698, 36111, 50521, 20370)
    # Cropped region is approximately square; set a readable final frame.
    shape.width = Inches(5.75)
    shape.height = Inches(5.75)

    doc.core_properties.title = "Feedback cần sửa VWork"
    doc.core_properties.subject = "Góp ý về thao tác gán vai trò"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()

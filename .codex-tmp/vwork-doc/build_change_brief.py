from __future__ import annotations

import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


SOURCE = Path(r"E:\2026\Cải tiến\VWork\VWork_feedback-mockup_code-UC-van-hanh-dao-tao\VWork_feedback-mockup_code-UC-van-hanh-dao-tao_07_can-xac-nhan.docx")
OUT = Path(r"E:\2026\Cải tiến\VWork\VWork_feedback-mockup_code-UC-van-hanh-dao-tao\VWork_feedback-change-brief_code-UC-van-hanh-dao-tao_09_da-trien-khai.docx")
TMP = Path(r"E:\tool\01. CODE\testVNB\.codex-tmp\vwork-doc\images-08")


def set_cell_fill(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if value and node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    elif not value and node is not None:
        p_pr.remove(node)


def extract_images_in_document_order(source: Path):
    TMP.mkdir(parents=True, exist_ok=True)
    doc = Document(source)
    ordered = []
    idx = 0
    for p in doc.paragraphs:
        for blip in p._p.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            if not rid:
                continue
            part = doc.part.related_parts[rid]
            ext = Path(part.partname).suffix or ".png"
            target = TMP / f"feedback-{idx + 1:02d}{ext}"
            target.write_bytes(part.blob)
            ordered.append(target)
            idx += 1
    if len(ordered) != 8:
        raise RuntimeError(f"Expected 8 images, found {len(ordered)}")
    return ordered


def add_label_paragraph(doc, label: str, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = False
    r = p.add_run(label + ": ")
    r.bold = True
    p.add_run(text)
    return p


def add_bullets(doc, label: str, items):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(7)
    h.paragraph_format.space_after = Pt(3)
    h.add_run(label).bold = True
    set_keep_with_next(h)
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_screenshot(doc, path: Path, number: int, width: float):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", f"Ảnh chụp giao diện VWork minh họa feedback, hình {number}")
    picture._inline.docPr.set("title", f"Bằng chứng giao diện VWork hình {number}")
    cap = doc.add_paragraph(f"Hình {number}. Bằng chứng giao diện do người dùng cung cấp")
    cap.style = doc.styles["Caption"]
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)


def add_feedback(doc, item, feedback_images, index, caption_start):
    doc.add_page_break()
    title = doc.add_paragraph(style="Heading 1")
    title.add_run(f"FB {index:02d}  {item['title']}")
    add_label_paragraph(doc, "Màn hình và vai trò", item["screen"])
    add_label_paragraph(doc, "Mức xác thực", item["evidence"])
    add_label_paragraph(doc, "Feedback gốc đã chuẩn hóa", item["source"])
    add_label_paragraph(doc, "Hiện trạng và nguyên nhân", item["current"])
    add_label_paragraph(doc, "Kết quả mong muốn", item["expected"])
    add_bullets(doc, "Acceptance criteria", item["ac"])
    add_label_paragraph(doc, "Phạm vi kỹ thuật dự kiến", item["scope"])
    add_bullets(doc, "Trạng thái cần kiểm tra", item["states"])
    add_label_paragraph(doc, "Rủi ro hồi quy", item["risk"])
    add_bullets(doc, "Cách kiểm tra sau khi sửa", item["test"])
    if item.get("decision"):
        add_label_paragraph(doc, "Quyết định đã xác nhận", item["decision"])
    if item.get("status"):
        add_label_paragraph(doc, "Trạng thái triển khai", item["status"])
    if item.get("verification"):
        add_bullets(doc, "Bằng chứng kiểm tra", item["verification"])
    for offset, image in enumerate(feedback_images):
        add_screenshot(doc, image, caption_start + offset, item.get("image_width", 6.25))


images = extract_images_in_document_order(SOURCE)

items = [
    {
        "title": "Làm nổi bật thao tác Gán vai trò",
        "screen": "Ekip khóa học; người dùng quản lý vận hành hoặc quản lý ekip.",
        "evidence": "Đã xác nhận bằng ảnh và code.",
        "source": "Vị trí hiện tại của thao tác Gán vai trò khó tìm; cần giúp người dùng nhận biết và thao tác nhanh hơn.",
        "current": "CourseControlPanel đặt hai ô chọn tài khoản và vai trò cùng nút Gán vai trò ở cuối danh sách ekip. Nút dùng kiểu mặc định, kích thước nhỏ và không có phân cấp thị giác rõ; vì vậy thao tác chính dễ bị bỏ sót.",
        "expected": "Thao tác Gán vai trò trở thành hành động chính, có nhãn rõ, trạng thái chọn hợp lệ và phản hồi thành công hoặc lỗi ngay trong ngữ cảnh Ekip khóa học.",
        "ac": [
            "Người có quyền nhìn thấy hành động Gán vai trò ngay khi mở khu vực Ekip khóa học, không phải dò ở cuối vùng nội dung.",
            "Nút chỉ khả dụng khi đã chọn tài khoản và vai trò hợp lệ; trạng thái disabled có thể nhận biết.",
            "Sau khi gán thành công, thành viên và vai trò mới xuất hiện trong danh sách mà không cần tải lại trang.",
            "Lỗi trùng vai trò, thiếu quyền hoặc lỗi API được thông báo tại chỗ; không làm mất lựa chọn đang nhập.",
            "Có thể thao tác bằng bàn phím và hiển thị đúng ở breakpoint hẹp.",
        ],
        "scope": "TrainingOperationsPage.tsx tại CourseControlPanel; TrainingOperationsPage.css; command gán vai trò và thông báo kết quả hiện có.",
        "states": ["Mặc định, chưa chọn tài khoản hoặc vai trò.", "Đang gửi, thành công, lỗi và không đủ quyền.", "Danh sách ekip dài; desktop và mobile."],
        "risk": "Di chuyển control có thể làm thay đổi bố cục modal khóa học hoặc tạo hai điểm gán vai trò nếu không loại bỏ vị trí cũ.",
        "test": ["Kiểm tra vai trò operations/admin/manager và người không có quyền.", "Gán mới, gán trùng, lỗi mạng và tải lại dữ liệu.", "Kiểm tra tab order, focus và bố cục ở chiều rộng nhỏ."],
        "decision": "Giữ thao tác ngay trong khu vực Ekip khóa học và làm nổi bật bằng một hàng phân công có nhãn rõ, CTA chính và bố cục responsive; không tạo modal mới.",
        "status": "PARTIAL - Đã triển khai local; còn cần kiểm tra trực quan sau đăng nhập trên test-vinabrain.",
        "verification": ["Architecture checks và production build đạt.", "Đã kiểm tra CTA có nhãn, disabled state và dùng pattern/token hiện có trong code."],
    },
    {
        "title": "Tách tiến độ checklist khỏi điều kiện Gửi duyệt",
        "screen": "Chi tiết công việc của thành viên ekip.",
        "evidence": "Đã xác nhận mâu thuẫn trực tiếp giữa feedback và UI/domain hiện tại.",
        "source": "Tích đủ checklist phải thể hiện 100%; gửi duyệt là hành vi riêng. Thiết kế lại: checkbox căn trái, ưu tiên không gian cho nội dung, cho phép tải ảnh minh chứng, bỏ Lưu bản nháp và không bắt buộc điền hoặc tải đủ mới được gửi duyệt.",
        "current": "UI chỉ cho Gửi duyệt khi toàn bộ checklist đã tích, mọi tiêu chí đã có minh chứng và đã có kết quả kèm minh chứng tổng. Domain submitReview cũng chặn nếu checklist chưa 100%, thiếu minh chứng từng tiêu chí hoặc thiếu minh chứng đầu ra. Tiến độ khi gửi duyệt bị đặt về 80%, nên không khớp yêu cầu 'tích hết là 100%'. Minh chứng từng tiêu chí hiện là ô nhập URL/ID/path, chưa có tải ảnh trực tiếp. Nút Lưu bản nháp đang tồn tại.",
        "expected": "Tiến độ phản ánh tỷ lệ checklist và đạt 100% khi tích đủ. Gửi duyệt là chuyển trạng thái độc lập, không bị ràng buộc phải hoàn tất hoặc upload đủ nếu nghiệp vụ cho phép. Mỗi tiêu chí có thể nhập nội dung và tải ảnh minh chứng; giao diện gọn, checkbox ở trái.",
        "ac": [
            "Tỷ lệ tiến độ được tính nhất quán từ checklist; tích toàn bộ hiển thị 100% trước và sau khi gửi duyệt.",
            "Gửi duyệt không bị khóa chỉ vì checklist, ghi chú hoặc minh chứng còn thiếu theo quy tắc được xác nhận.",
            "Bỏ nút Lưu bản nháp; thay đổi checklist và nội dung được lưu theo cơ chế rõ ràng, không làm mất dữ liệu khi đóng/mở lại.",
            "Mỗi tiêu chí cho phép đính kèm ảnh bằng API upload hiện có, hiển thị tên/preview và cho phép xóa trước khi gửi.",
            "Checkbox nằm bên trái nhãn; nội dung và minh chứng có đủ chiều rộng, không tràn hoặc che nút.",
            "Lỗi upload/lưu/gửi duyệt giữ lại dữ liệu người dùng và nêu cách thử lại.",
        ],
        "scope": "TaskDrawer trong TrainingOperationsPage.tsx và CSS; domain updateTaskProgress/submitReview; API file; test domain và luồng UI.",
        "states": ["0%, đang làm, 100%, gửi duyệt và làm lại.", "Không có minh chứng, đang upload, upload lỗi, xóa ảnh.", "Đang lưu nền, lưu thành công, lưu lỗi, mất mạng và mở lại drawer."],
        "risk": "Nới điều kiện submit ở UI mà không sửa domain sẽ vẫn lỗi. Sửa domain ảnh hưởng toàn bộ tính toàn vẹn phiếu duyệt và snapshot; autosave có thể tạo race condition nếu người dùng tích nhanh hoặc đóng drawer.",
        "test": ["Gửi duyệt tại 0%, một phần và 100% theo quyết định nghiệp vụ.", "Tải ảnh hợp lệ/không hợp lệ, nhiều ảnh, lỗi mạng và xóa ảnh.", "Đóng/mở lại để xác nhận dữ liệu đã lưu; kiểm tra snapshot reviewer nhận được.", "Xác minh tiến độ không tụt từ 100% về 80% khi chuyển IN_REVIEW."],
        "decision": "Cho phép Gửi duyệt ở mọi mức tiến độ, kể cả 0% và không có minh chứng. Reviewer quyết định Đạt hoặc trả lại.",
        "status": "PARTIAL - Nghiệp vụ và UI đã triển khai; upload và autosave cần smoke test sau đăng nhập trên test-vinabrain.",
        "verification": ["Domain test chứng minh gửi duyệt ở 0%, snapshot output rỗng hợp lệ, REWORK giữ 0% và PASS đưa tiến độ về 100%.", "Architecture checks, typecheck và production build đạt."],
        "image_width": 5.35,
    },
    {
        "title": "Hiển thị chi tiết checklist trong màn hình nghiệm thu",
        "screen": "Hàng đợi nghiệm thu của quản lý ekip.",
        "evidence": "Đã xác nhận bằng ảnh và code.",
        "source": "Quản lý ekip cần xem được chi tiết checklist công việc khi mở chi tiết.",
        "current": "ReviewQueue hiện chỉ hiển thị tên công việc, kết quả, minh chứng tổng, các chỉ số tóm tắt và nút Đạt/Yêu cầu làm lại. Checklist và minh chứng theo từng tiêu chí đã có trong submission snapshot nhưng chưa được render cho reviewer.",
        "expected": "Khi mở chi tiết, reviewer thấy đúng snapshot của lần nộp: từng tiêu chí, trạng thái tick, nội dung/minh chứng, kết quả đầu ra và lịch sử liên quan trước khi ra quyết định.",
        "ac": [
            "Nút hoặc vùng Chi tiết mở được đầy đủ checklist của đúng submission đang chờ duyệt.",
            "Mỗi tiêu chí hiển thị trạng thái hoàn thành và danh sách minh chứng có thể mở/tải theo quyền.",
            "Reviewer nhìn rõ dữ liệu snapshot của lần nộp, không bị thay đổi bởi chỉnh sửa sau đó.",
            "Trạng thái không có checklist/minh chứng được thể hiện rõ, không để vùng trống gây hiểu nhầm.",
            "Đạt và Yêu cầu làm lại vẫn thao tác được sau khi xem chi tiết trên desktop và màn hình hẹp.",
        ],
        "scope": "ReviewQueue; component chi tiết submission dùng lại dữ liệu task.submissions/taskSnapshot; kiểm soát quyền mở file.",
        "states": ["Có/không có checklist; tiêu chí đạt/chưa đạt.", "Minh chứng hợp lệ, thiếu, không truy cập được.", "Loading, error, permission denied và lịch sử nhiều lần nộp."],
        "risk": "Nếu đọc trực tiếp task hiện tại thay vì snapshot, reviewer có thể đánh giá nhầm dữ liệu đã thay đổi sau lúc gửi.",
        "test": ["Nộp hai phiên bản khác nhau và xác nhận chi tiết từng snapshot.", "Mở từng minh chứng với đúng quyền.", "Kiểm tra empty/error và hành động review sau khi mở chi tiết."],
        "status": "PARTIAL - Đã render snapshot checklist và minh chứng trong code; còn cần smoke test quyền mở file trên test-vinabrain.",
        "verification": ["Architecture checks xác nhận reviewer có Xem chi tiết và đọc checklist snapshot.", "Domain test xác nhận snapshot của submission được giữ bất biến."],
    },
    {
        "title": "Sửa lỗi layout khi mở chọn Người thực hiện",
        "screen": "Danh sách công việc và drawer chi tiết; người phân công công việc.",
        "evidence": "Ảnh xác nhận lỗi; code cho thấy dialog chưa dùng overlay portal sẵn có. Cần tái hiện runtime để chốt root cause.",
        "source": "Khi nhấn vào Người thực hiện, khung giao diện bị tràn hoặc chồng lớp như ảnh.",
        "current": "AssigneePickerDialog dùng backdrop position fixed nhưng được render bên trong cây layered-task-browser thay vì helper overlay portal lên document.body. Trong bố cục bảng cộng drawer, stacking/containing context có thể làm dialog bị cắt, lệch hoặc đè lên panel. Đây là nguyên nhân dự kiến, chưa phải kết luận runtime.",
        "expected": "Dialog chọn người thực hiện luôn nằm trên toàn bộ workspace, căn giữa vùng nhìn, cuộn nội dung bên trong và không làm dịch/chèn bảng hoặc drawer.",
        "ac": [
            "Mở dialog không làm thay đổi kích thước hay vị trí bảng và drawer bên dưới.",
            "Backdrop phủ đúng viewport; dialog không bị cắt ở cạnh phải/dưới và có max-height phù hợp.",
            "Có thể đóng bằng nút, Escape và click backdrop theo pattern hiện có; focus được đưa vào dialog và trả về control mở.",
            "Danh sách dài cuộn trong dialog; thanh tìm kiếm và hành động chính vẫn nhìn thấy.",
            "Không tạo lỗi stacking với modal khóa học hoặc shell PeopleOne.",
        ],
        "scope": "LayeredTasks, AssigneePickerDialog, renderTrainingOverlay và các selector assignment-dialog/task-layout trong CSS.",
        "states": ["Danh sách ngắn/dài, không có kết quả, loading và error.", "Drawer mở/đóng; viewport 1366x768, 1920x1080 và breakpoint hẹp.", "Keyboard focus và reduced motion nếu có."],
        "risk": "Portal có thể tách dialog khỏi context CSS hoặc event boundary; cần giữ theme variables và cleanup body/focus đúng cách.",
        "test": ["Tái hiện đúng viewport trong ảnh trước sửa.", "Kiểm tra bounding box dialog/backdrop và không có horizontal scroll.", "Tab/Escape/focus return; chọn người, hủy, tìm không có kết quả."],
        "status": "PARTIAL - Đã portal dialog lên document.body, giới hạn chiều cao, cuộn nội dung và trap focus; chưa tái hiện sau đăng nhập.",
        "verification": ["Architecture checks xác nhận portal, Tab trap và vùng body cuộn độc lập.", "Typecheck và production build đạt."],
    },
    {
        "title": "Đồng bộ tài khoản Quản lý vận hành vào danh sách phân công",
        "screen": "Quản lý vận hành và dialog giao việc; operations/manager.",
        "evidence": "Đã xác nhận cơ chế lọc trong code; chưa có dữ liệu runtime để xác định tài khoản thiếu ở bước đồng bộ, liên kết profile, trạng thái active hay role.",
        "source": "Không tìm thấy các tài khoản đã tạo ở Quản lý vận hành khi chọn người thực hiện; ví dụ nhập tên hoặc email Chiêu Anh không ra.",
        "current": "Frontend tìm trên directory được API trả về. API chỉ đưa vào directory tài khoản có profile liên kết, active và có role manager/member; tài khoản chỉ tồn tại ở vplanning_users hoặc chưa được reconcile profile/role sẽ có thể xuất hiện ở accountDirectory nhưng không được phép nhận việc. Vì vậy lỗi không nằm ở phép tìm kiếm đơn thuần.",
        "expected": "Mọi tài khoản active được tạo đúng quy trình Quản lý vận hành đều xuất hiện trong tìm kiếm phân công, không phụ thuộc role; tài khoản inactive hoặc chưa liên kết profile không được nhận việc.",
        "ac": [
            "Sau khi tạo/kích hoạt tài khoản, tìm bằng tên không dấu/có dấu và email trả đúng người theo quy tắc tìm kiếm thống nhất.",
            "Tài khoản thiếu profile hoặc inactive được hiển thị trạng thái không thể chọn và lý do cho người có quyền quản lý tài khoản, hoặc được reconcile tự động theo contract đã chọn.",
            "Server vẫn kiểm tra assignable và phạm vi project/course/class; không chỉ dựa vào trạng thái disabled ở UI.",
            "Sau khi cập nhật role hoặc active, danh sách được refresh/invalidate và không cần đăng xuất.",
            "Không hiển thị tài khoản ngoài tenant hoặc ngoài phạm vi được phép.",
        ],
        "scope": "loadVWorkAccountDirectory và normalizeAssignmentCommand; API quản lý tài khoản/reconciliation; loadWorkspace và AssigneePickerDialog.",
        "states": ["Đã liên kết/chưa liên kết profile; active/inactive.", "Có role member/manager, role khác hoặc nhiều role.", "Đúng/sai project, course, class; dữ liệu vừa tạo và dữ liệu cũ."],
        "risk": "Cho tất cả tài khoản active vào danh sách nhưng bỏ kiểm tra server về profile, trạng thái hoặc scope có thể cấp việc sai tenant. Thay đổi chỉ gỡ điều kiện role, không gỡ các ranh giới còn lại.",
        "test": ["Tạo tài khoản mới rồi tìm ngay bằng tên/email; refresh và đăng nhập lại.", "Ma trận profile-linked, active, role và scope; xác minh thông báo lý do.", "Thử gọi API phân công trực tiếp với tài khoản không hợp lệ để bảo đảm server từ chối."],
        "decision": "Mọi tài khoản active được nhận việc, không giới hạn role member hoặc manager. Server vẫn giữ kiểm tra trạng thái active, profile liên kết và phạm vi project/course/class.",
        "status": "PARTIAL - Điều kiện role đã được gỡ ở API/UI và đã có API test; cần xác minh tài khoản thực tế Chiêu Anh trên test-vinabrain.",
        "verification": ["API test bổ sung tài khoản active chỉ có role nội dung và xác nhận vẫn assignable.", "API vẫn từ chối tài khoản inactive và tài khoản ngoài scope."],
    },
    {
        "title": "Rút gọn bố cục checklist trong drawer công việc",
        "screen": "Drawer chi tiết công việc của thành viên ekip.",
        "evidence": "Ảnh xác nhận khoảng trống lớn; code cho thấy mỗi tiêu chí và vùng minh chứng đang xếp dọc. Cách hiểu 'trên cùng một dòng' cần xác nhận.",
        "source": "Khi khung đủ rộng, các cụm như Đúng tên lớp và Đúng thời gian chạy lớp cần được bố trí trên cùng một dòng để tận dụng không gian.",
        "current": "Mỗi check-evidence-row là một khối grid xếp dọc, trong khi drawer có chiều rộng lớn. Nhãn ngắn vẫn chiếm một hàng riêng và vùng minh chứng tạo nhiều khoảng trắng, làm danh sách dài hơn cần thiết.",
        "expected": "Checklist dùng bố cục responsive: desktop tận dụng chiều ngang, nhãn và control thẳng hàng; màn hình hẹp tự xếp dọc mà không giảm khả năng đọc hoặc vùng bấm.",
        "ac": [
            "Ở desktop, các tiêu chí ngắn được bố trí theo lưới hai cột hoặc theo cấu trúc một hàng đã xác nhận.",
            "Checkbox luôn ở trái nhãn, vùng bấm đủ lớn và nhãn không bị cắt.",
            "Nội dung/minh chứng liên kết rõ với đúng tiêu chí, không gây nhầm khi hai tiêu chí cùng hàng.",
            "Ở breakpoint hẹp, các tiêu chí xếp một cột theo thứ tự dữ liệu.",
            "Không xuất hiện horizontal scroll hoặc khoảng trắng bất thường với nhãn dài.",
        ],
        "scope": "TaskDrawer markup và CSS check-evidence-row/check-row; dùng design token/pattern hiện có.",
        "states": ["Nhãn ngắn/dài, có/không có minh chứng.", "Desktop, tablet và mobile.", "Checked, unchecked, disabled và focus."],
        "risk": "Ghép hai tiêu chí vào cùng hàng có thể làm người dùng liên kết nhầm minh chứng; cần grouping và responsive rõ.",
        "test": ["So sánh chiều cao drawer trước/sau với 2, 6 và 12 tiêu chí.", "Kiểm tra 320/768/1366/1920 px và nhãn dài.", "Tab order, click label để đổi checkbox và upload đúng tiêu chí."],
        "decision": "Dùng bố cục chuyên nghiệp và cân đối: mỗi tiêu chí là một nhóm rõ ràng; lưới hai cột chỉ kích hoạt khi drawer đủ rộng, tự về một cột ở màn hình hẹp.",
        "status": "PARTIAL - Đã triển khai responsive container query; còn cần đối chiếu trực quan ở các viewport mục tiêu.",
        "verification": ["Architecture checks xác nhận lưới hai cột tại container từ 720 px.", "Checkbox luôn nằm bên trái nhãn và mỗi minh chứng nằm trong đúng card tiêu chí."],
        "image_width": 5.45,
    },
    {
        "title": "Bổ sung ghi chú khi trả lại công việc",
        "screen": "Hàng đợi nghiệm thu của quản lý ekip và thông báo của thành viên ekip.",
        "evidence": "Đã xác nhận bằng code: domain đã bắt buộc comment cho REWORK nhưng UI đang gửi một câu cố định.",
        "source": "Bổ sung phần note review của quản lý ekip khi trả lại cho thành viên ekip.",
        "current": "ReviewQueue không có ô nhập ghi chú. Khi bấm Yêu cầu làm lại, UI tự gửi câu 'Cần bổ sung cấu hình hoặc minh chứng theo tiêu chí xác nhận.' Domain đã hỗ trợ và bắt buộc comment khi REWORK; comment được lưu trong reviews, notification và audit.",
        "expected": "Reviewer nhập lý do cụ thể trước khi trả lại. Ghi chú được lưu đúng lần review, hiển thị cho thành viên ở công việc/thông báo và còn trong lịch sử khi nộp lại.",
        "ac": [
            "Yêu cầu làm lại mở vùng nhập ghi chú; không cho xác nhận khi ghi chú rỗng hoặc chỉ có khoảng trắng.",
            "Ghi chú được gửi nguyên vẹn, gắn với submission/review hiện tại và hiển thị cho đúng assignee.",
            "Thành viên nhìn thấy ghi chú ngay khi mở lại công việc và từ thông báo liên quan.",
            "Lịch sử giữ riêng ghi chú của từng vòng rework/resubmit.",
            "Lỗi gửi không làm mất nội dung reviewer đã nhập; tránh tạo review trùng khi bấm nhiều lần.",
        ],
        "scope": "ReviewQueue; domain reviewTask hiện có; hiển thị reviews/notifications trong TaskDrawer và lịch sử; test vòng rework.",
        "states": ["Chưa nhập, đang nhập, quá dài, đang gửi, lỗi và thành công.", "Nhiều vòng trả lại; reviewer/assignee khác nhau.", "Notification đã đọc/chưa đọc."],
        "risk": "Nếu chỉ thêm input nhưng vẫn dùng comment cố định hoặc không gắn đúng submission, lịch sử sẽ sai. Cần giới hạn/escape nội dung để tránh lỗi hiển thị.",
        "test": ["Trả lại với ghi chú rỗng/hợp lệ; giả lập lỗi mạng và double click.", "Đăng nhập thành viên để kiểm tra notification và chi tiết task.", "Nộp lại rồi trả lại lần hai, xác nhận lịch sử hai ghi chú tách biệt."],
        "status": "PARTIAL - Ô ghi chú bắt buộc và hiển thị feedback cho thành viên đã triển khai; notification runtime chưa được smoke test.",
        "verification": ["Domain test xác nhận comment rỗng bị từ chối và comment hợp lệ được lưu đúng review.", "Architecture checks xác nhận UI không còn gửi câu cố định và có vùng nhập lý do."],
    },
    {
        "title": "Hoàn thiện luồng end to end vận hành đào tạo",
        "screen": "Toàn bộ UC vận hành đào tạo từ quản lý tài khoản đến nghiệm thu.",
        "evidence": "Yêu cầu phạm vi tổng thể; các command nền đã tồn tại nhưng các feedback FB 01 đến FB 07 cho thấy luồng người dùng chưa khép kín.",
        "source": "Không xử lý rời từng màn hình; toàn bộ logic phải chạy được end to end.",
        "current": "Domain có các bước phân công, cập nhật checklist, nộp output, gửi duyệt, review, notification và audit. Tuy nhiên UI và quy tắc còn đứt đoạn: tài khoản có thể không vào directory, submit bị ràng buộc khác feedback, reviewer thiếu checklist/ghi chú, và layout chọn người có lỗi. Test hiện tại xác nhận kiến trúc và contract nền chứ chưa chứng minh hành trình UI hoàn chỉnh trên dữ liệu runtime.",
        "expected": "Một lớp học thật có thể đi liền mạch qua các vai trò: tạo/đồng bộ tài khoản, gán ekip, giao việc, thực hiện checklist và minh chứng, gửi duyệt, xem chi tiết, trả lại kèm ghi chú, sửa/nộp lại, duyệt đạt và hoàn thành 100%.",
        "ac": [
            "Tài khoản hợp lệ xuất hiện trong Ekip khóa học và danh sách giao việc theo đúng quyền/scope.",
            "Gán vai trò và giao việc cập nhật tức thời; assignee nhìn thấy đúng task.",
            "Thành viên cập nhật checklist/minh chứng, đóng mở lại không mất dữ liệu và gửi duyệt theo quy tắc FB 02 đã xác nhận.",
            "Reviewer xem đúng checklist snapshot và minh chứng, có thể trả lại với ghi chú bắt buộc.",
            "Thành viên nhận được ghi chú, sửa và nộp lại; lịch sử các vòng không bị ghi đè.",
            "Khi reviewer chọn Đạt, task ở DONE, progress 100%, completedAt/review/notification/audit nhất quán.",
            "Refresh hoặc đăng nhập lại ở từng vai trò không làm mất hoặc lộ dữ liệu ngoài phạm vi.",
            "Luồng lỗi API/upload/concurrency có thông báo và không tạo dữ liệu trùng.",
        ],
        "scope": "Các component UC vận hành đào tạo; domain commands; API state, user và file; persistence, notification, audit và bộ test E2E nhiều vai trò.",
        "states": ["Operations, manager/reviewer và member/assignee.", "Happy path, rework một/nhiều vòng, lỗi upload/API và concurrent update.", "Dữ liệu cũ/mới; reload/login lại; desktop và breakpoint liên quan."],
        "risk": "Sửa từng điểm mà không khóa acceptance criteria xuyên vai trò có thể tạo trạng thái UI đúng cục bộ nhưng domain/API không nhất quán. Dữ liệu runtime sai tenant/lớp có thể làm kết quả kiểm thử gây hiểu nhầm.",
        "test": ["Tạo fixture lớp/course riêng và ba tài khoản đúng vai trò; ghi lại ID môi trường.", "Chạy kịch bản xuyên suốt từ provisioning đến PASS, sau đó lặp lại nhánh REWORK và resubmit.", "Đối chiếu state lưu trữ, submission snapshot, review, notification và audit ở mỗi bước.", "Chạy lại bộ test architecture/domain/API/user/file và thêm browser E2E cho hành trình chính."],
        "decision": "Dùng các link thuộc môi trường test-vinabrain để chạy E2E; không dùng dữ liệu production.",
        "status": "PARTIAL - Chu trình domain/API đã được kiểm tra xuyên suốt; browser E2E trên test-vinabrain chưa chạy do phiên local dừng ở màn đăng nhập.",
        "verification": ["Kịch bản tự động đã chạy 0% -> gửi duyệt -> REWORK có lý do -> gửi lại -> PASS 100%.", "Bảy bộ test VWork, typecheck và production build đều đạt."],
    },
]

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9.4)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.10

for style_name, size, before, after in (("Title", 22, 0, 12), ("Heading 1", 15, 10, 6), ("Heading 2", 12, 8, 4)):
    st = styles[style_name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

styles["Caption"].font.name = "Arial"
styles["Caption"].font.size = Pt(8.5)
styles["Caption"].font.color.rgb = RGBColor(89, 89, 89)
styles["Caption"].font.italic = True

# Word's built-in Title style may carry a theme border in some renderers.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

title = doc.add_paragraph(style="Title")
title.add_run("Change brief feedback VWork vận hành đào tạo")
title_direct_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
if title_direct_border is not None:
    title._p.get_or_add_pPr().remove(title_direct_border)
subtitle = doc.add_paragraph("Phạm vi: UC vận hành đào tạo | Nguồn: file feedback cập nhật ngày 09 tháng 09 năm 2026")
subtitle.paragraph_format.space_after = Pt(12)
subtitle.runs[0].italic = True
subtitle.runs[0].font.color.rgb = RGBColor(80, 80, 80)

intro = doc.add_paragraph()
intro.add_run("Kết luận: ").bold = True
intro.add_run("Tài liệu ghi nhận 8 feedback độc lập và cập nhật kết quả triển khai local sau khi đã chốt các quyết định nghiệp vụ. Cả 8 mục đã được xử lý trong code và qua kiểm tra tự động; kiểm tra trực quan, upload và browser E2E trên test-vinabrain vẫn cần một phiên đăng nhập hợp lệ trước khi nghiệm thu cuối.")

doc.add_paragraph("Tóm tắt đối chiếu", style="Heading 1")
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(1.45), Cm(7.5), Cm(3.4), Cm(4.5)]
headers = ["Mã", "Feedback", "Kết quả hiện tại", "Quyết định"]
for i, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
    cell.width = width
    set_cell_fill(cell, "1F4E78")
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(header)
    r.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
set_repeat_table_header(table.rows[0])

summary_rows = [
    ("FB 01", items[0]["title"], "PARTIAL", "Inline CTA rõ ràng"),
    ("FB 02", items[1]["title"], "PARTIAL", "Cho gửi ở 0%"),
    ("FB 03", items[2]["title"], "PARTIAL", "Dùng snapshot"),
    ("FB 04", items[3]["title"], "PARTIAL", "Portal toàn viewport"),
    ("FB 05", items[4]["title"], "PARTIAL", "Mọi tài khoản active"),
    ("FB 06", items[5]["title"], "PARTIAL", "Lưới responsive"),
    ("FB 07", items[6]["title"], "PARTIAL", "Bắt buộc ghi chú"),
    ("FB 08", items[7]["title"], "PARTIAL", "test-vinabrain"),
]
for ridx, row in enumerate(summary_rows, 1):
    cells = table.add_row().cells
    for cidx, (cell, value, width) in enumerate(zip(cells, row, widths)):
        cell.width = width
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ridx % 2 == 0:
            set_cell_fill(cell, "EDF3F8")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cidx in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(value)
set_table_borders(table)

doc.add_paragraph("Bằng chứng kỹ thuật đã kiểm tra", style="Heading 1")
add_bullets(doc, "Phạm vi đối chiếu", [
    "Frontend: TrainingOperationsPage.tsx và TrainingOperationsPage.css.",
    "Domain: cập nhật tiến độ, nộp output, gửi duyệt, review, notification và audit.",
    "API: state, directory tài khoản, provisioning/reconciliation và upload file.",
    "Bảy bộ kiểm tra: architecture, domain UC01-UC16 và E2E review, API auth/persistence/concurrency, user provisioning, upload metadata, multi-role và login reconciliation đều đạt.",
    "Typecheck và production build đạt; build chỉ còn cảnh báo kích thước chunk đã tồn tại ở phạm vi toàn ứng dụng.",
])
add_label_paragraph(doc, "Giới hạn", "Browser E2E chưa chạy vì phiên local yêu cầu đăng nhập. Các file design context không tồn tại tại đường dẫn quy định; phần triển khai giữ PeopleOne làm product shell theo code hiện tại.")

image_groups = [
    [images[0]],
    [images[1]],
    [images[2]],
    [images[3], images[4]],
    [images[5]],
    [images[6]],
    [images[7]],
    [],
]
caption_number = 1
for i, (item, feedback_images) in enumerate(zip(items, image_groups), 1):
    add_feedback(doc, item, feedback_images, i, caption_number)
    caption_number += len(feedback_images)

doc.add_page_break()
doc.add_paragraph("Kịch bản nghiệm thu end to end đề xuất", style="Heading 1")
steps = [
    "Operations tạo hoặc kích hoạt tài khoản thành viên và quản lý; xác nhận profile, role và scope đồng bộ.",
    "Quản lý mở khóa học, tìm thấy tài khoản và gán vai trò trong Ekip khóa học.",
    "Quản lý chọn công việc, mở dialog Người thực hiện, tìm đúng tài khoản và giao việc.",
    "Thành viên mở task, cập nhật checklist/nội dung/minh chứng; đóng mở lại không mất dữ liệu.",
    "Thành viên gửi duyệt theo quy tắc FB 02 đã chốt; hệ thống tạo snapshot và thông báo reviewer.",
    "Reviewer mở chi tiết, xem checklist/minh chứng của snapshot và chọn Yêu cầu làm lại kèm ghi chú.",
    "Thành viên nhận ghi chú, sửa nội dung, nộp lại; lịch sử vòng trước được giữ nguyên.",
    "Reviewer chọn Đạt; task DONE và progress 100%, notification/audit/persistence nhất quán sau refresh.",
]
for idx, step in enumerate(steps, 1):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(step)

doc.add_paragraph("Điều kiện nghiệm thu còn lại", style="Heading 1")
add_bullets(doc, "Cần kiểm tra trên test-vinabrain", [
    "Đăng nhập bằng tài khoản operations, manager và member thuộc đúng tenant/course test.",
    "Chạy browser E2E cho giao việc, upload, gửi duyệt, xem snapshot, REWORK và PASS.",
    "Đối chiếu trực quan dialog, checklist ở desktop/mobile và xác minh tài khoản Chiêu Anh xuất hiện trong tìm kiếm.",
    "Xác nhận notification, persistence sau refresh và quyền mở file ở từng vai trò.",
])
add_label_paragraph(doc, "GO hoặc HOLD", "GO cho phạm vi code local và kiểm tra tự động. HOLD cho nghiệm thu cuối hoặc triển khai production cho đến khi browser E2E trên test-vinabrain đạt. Không deploy, commit, push hoặc thay đổi dữ liệu production trong phạm vi công việc này.")

for section in doc.sections:
    header = section.header.paragraphs[0]
    header.text = "VWork  Change brief UC vận hành đào tạo"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)
    footer = section.footer.paragraphs[0]
    footer.text = "Tài liệu làm rõ phạm vi và tiêu chí nghiệm thu"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(110, 110, 110)

doc.core_properties.title = "Change brief feedback VWork vận hành đào tạo"
doc.core_properties.subject = "Chuẩn hóa 8 feedback và tiêu chí nghiệm thu"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("DOCX_CREATED")

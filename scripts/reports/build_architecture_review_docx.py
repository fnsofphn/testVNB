from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "reports" / "Bao-cao-review-kien-truc-3000-nguoi-dung.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def format_table(table, widths=None, header_fill="1F4E79"):
    table.style = "Table Grid"
    if widths:
        set_table_width(table, widths)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
                    run.font.size = Pt(9.2)
        if row_idx == 0:
            for cell in row.cells:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        elif row_idx % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, "F3F7FA")


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        set_run_font(run, size=15 if level == 1 else 12, bold=True, color="000000")
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=10.4, bold=True)
        run = paragraph.add_run(text)
    else:
        run = paragraph.add_run(text)
    set_run_font(run, size=10.4)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=10.1)


def add_table(doc, headers, rows, widths=None, header_fill="1F4E79"):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for idx, value in enumerate(headers):
        hdr[idx].text = value
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    format_table(table, widths=widths, header_fill=header_fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.4)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.color.rgb = RGBColor(0, 0, 0)


def add_title_page(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Báo cáo review kiến trúc phần mềm")
    set_run_font(run, size=24, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Mục tiêu đáp ứng tối thiểu 3.000 người dùng truy cập đồng thời")
    set_run_font(run, size=13, bold=True, color="333333")

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(16)
    run = meta.add_run("Hệ thống: VContent / Vinabrain | Phạm vi: Kiến trúc tổng thể, VLearning, VTraining, VSurvey, VEvent | Ngày lập: 14/09/2026")
    set_run_font(run, size=9.8, color="555555")

    add_body(
        doc,
        "Tài liệu này tổng hợp kết quả review ở mức Solution Architecture dựa trên cấu trúc mã nguồn hiện tại. Kết luận chính là hệ thống đã có nền tảng phù hợp cho quy mô vừa, nhưng để vận hành ổn định ở mốc 3.000 người dùng đồng thời cần tăng cường cache phân tán, rate limit phân tán, hàng đợi xử lý nền, read model cho báo cáo, kiểm thử tải và observability sản xuất.",
    )
    add_body(
        doc,
        "Mô hình hiện tại có thể xem là React/Vite SPA triển khai trên Vercel, dùng Supabase làm backend dữ liệu chính, kèm một số API serverless cho các luồng nhạy cảm. Đây là hướng modular monolith thiên về database-centric backend, chưa phải microservices.",
        bold_lead="Nhận định tổng quan: ",
    )


def build_report():
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    add_heading(doc, "1 Tóm tắt điều hành", 1)
    add_table(
        doc,
        ["Hạng mục", "Đánh giá hiện tại", "Khuyến nghị chính"],
        [
            ["Mô hình kiến trúc", "React/Vite SPA + Vercel Serverless + Supabase. Frontend được chia module, backend chủ yếu là Supabase/RPC/API cục bộ.", "Giữ modular monolith ngắn hạn, bổ sung BFF/API contract cho write path và hot path."],
            ["Mục tiêu tải", "Cấu hình hiện có đặt VITE_CONCURRENCY_TARGET=1000, chưa thấy gate sản xuất cho 3.000 concurrent users.", "Nâng target lên 3.000, chạy load test theo module và đặt SLO trước release."],
            ["Cache và rate limit", "Có React Query cache và helper Redis/Upstash, nhưng rate limit mặc định local.", "Bật Redis/Upstash enforce cho production, thêm cache server-side cho read path nóng."],
            ["Database", "Supabase/Postgres là trung tâm; nhiều luồng client gọi trực tiếp bảng/RPC.", "Tối ưu index, connection pooler, read model/materialized view và giảm query trực tiếp từ client."],
            ["Vận hành", "Có telemetry/RUM ban đầu, có nhiều script test; monitoring/alerting sản xuất chưa thể xác nhận đầy đủ.", "Chuẩn hóa logging, tracing, Sentry/APM, dashboard DB/API, alert và runbook incident."],
        ],
        widths=[Inches(1.35), Inches(2.55), Inches(2.95)],
    )

    add_heading(doc, "2 Kiến trúc hiện tại", 1)
    add_body(doc, "Frontend sử dụng React 19, Vite, React Router và React Query. Ứng dụng được chia route theo các sản phẩm VLearning, VTraining, VSurvey, VEvent, VDiscussion, VWork và các module hỗ trợ.")
    add_body(doc, "Hosting dùng Vercel với region sin1, output static dist, các API JavaScript trong thư mục api chạy dạng serverless function. Static assets trong /assets đã có Cache-Control immutable, index.html no-cache.")
    add_body(doc, "Data layer dùng Supabase client ở frontend cho nhiều luồng đọc/ghi. Một số luồng nhạy cảm như VWork Training Operations đi qua API serverless với service role, kiểm tra token, phân quyền, optimistic concurrency, idempotency và audit.")

    add_table(
        doc,
        ["Thành phần", "Hiện trạng quan sát được", "Ý nghĩa với 3.000 concurrent users"],
        [
            ["Frontend SPA", "Route-level lazy loading đã được dùng.", "Tốt cho bundle ban đầu, nhưng cần kiểm soát kích thước route và preload tài nguyên."],
            ["API serverless", "Có API cho VWork Training Operations, file upload, user operation, rate limit/body limit helpers.", "Cần mở rộng vai trò BFF cho public write path, survey/event, export/report."],
            ["Supabase", "Client gọi trực tiếp nhiều bảng; RPC được dùng cho quiz, VLearning, VDiscussion, VEvent.", "Database dễ thành nút cổ chai nếu polling, report và write-heavy không được gom/batch/cache."],
            ["Realtime/polling", "Realtime nhiều nơi tắt, hybrid hoặc polling; VEvent polling 5s.", "Polling 5s với 3.000 người dùng có thể tạo khoảng 600 req/s chỉ cho một luồng."],
            ["Distributed state", "Đã có helper Redis/Upstash và atomic counter.", "Cần bật enforce trong production để rate limit/counter hoạt động xuyên instance."],
        ],
        widths=[Inches(1.45), Inches(2.75), Inches(2.65)],
    )

    add_heading(doc, "3 Rủi ro kiến trúc", 1)
    add_table(
        doc,
        ["Mức độ", "Rủi ro", "Tác động", "Hành động ưu tiên"],
        [
            ["Critical", "Chưa có bằng chứng capacity thật cho 3.000 concurrent users.", "Không thể cam kết SLA khi triển khai lớp/sự kiện lớn.", "Chạy k6/load test theo module, lưu baseline p95, error rate, RPS, DB load."],
            ["Critical", "Nhiều read/write trực tiếp từ client vào Supabase.", "Khó kiểm soát rate, transaction, idempotency và payload lớn.", "Đưa public write path và hot write path qua API/RPC/BFF."],
            ["High", "Rate limit mặc định local trên serverless.", "Không bảo vệ toàn hệ khi Vercel scale nhiều instance.", "Bật Redis/Upstash enforce và failure mode closed cho route nhạy cảm."],
            ["High", "Polling có thể tạo bão request.", "DB/PostgREST bị quá tải khi nhiều người cùng mở lớp/sự kiện.", "Adaptive polling, jitter, cache, realtime có kiểm soát cho presenter/monitor."],
            ["High", "Report/export xử lý phía client và đọc toàn bộ dữ liệu.", "Treo trình duyệt, timeout API, query nặng.", "Chuyển export/report sang background job và file download link."],
            ["Medium", "Chưa thấy monitoring stack sản xuất đầy đủ.", "Khó phát hiện sớm suy giảm hiệu năng hoặc lỗi DB.", "Thiết lập APM, Sentry, Supabase query insights, alert và runbook."],
            ["Medium", "Media/SCORM/file cần chính sách CDN rõ hơn.", "Tăng băng thông app, tải chậm, chi phí không kiểm soát.", "Chuẩn hóa object storage, CDN, signed URL, cache headers."],
        ],
        widths=[Inches(0.85), Inches(2.05), Inches(1.95), Inches(1.95)],
        header_fill="404040",
    )

    add_heading(doc, "4 Kiến trúc đề xuất", 1)
    add_body(doc, "Trong ngắn hạn không cần chuyển ngay sang microservices. Hướng phù hợp là modular monolith có backend contract rõ ràng, dùng Supabase/Postgres làm nguồn dữ liệu chính nhưng giảm truy cập trực tiếp từ client ở các luồng chịu tải cao hoặc cần transaction.")
    add_body(doc, "Luồng đề xuất: Client SPA -> CDN/Vercel Edge -> API Gateway hoặc BFF Serverless -> Supabase/Postgres -> Redis, Queue, Object Storage, CDN và Monitoring.")
    add_table(
        doc,
        ["Thành phần bổ sung", "Vai trò", "Ưu tiên"],
        [
            ["Redis/Upstash", "Distributed rate limit, short cache, counters, idempotency guard.", "P0"],
            ["Queue/Worker", "Email/SMS, report/export, certificate, scoring, import/sync dữ liệu.", "P0"],
            ["Object Storage + CDN", "Video, SCORM, tài liệu, ảnh, file minh chứng.", "P0"],
            ["Read model/materialized view", "Dashboard, survey summary, event stats, completion report.", "P1"],
            ["Monitoring stack", "APM, error tracking, tracing, log correlation, DB/API alerts.", "P0"],
            ["Connection pooling", "Giảm áp lực kết nối Postgres khi traffic tăng.", "P0"],
            ["Read replica/analytics DB", "Báo cáo lớn, phân tích lịch sử, export quy mô cao.", "P2"],
        ],
        widths=[Inches(1.65), Inches(3.75), Inches(0.85)],
    )

    add_heading(doc, "5 Review theo module", 1)
    add_heading(doc, "5.1 VLearning", 2)
    add_body(doc, "VLearning đã có hướng tốt khi dùng RPC cho learner state, SCORM runtime và cơ chế buffer/retry phía client. Tuy nhiên với 3.000 người học, cần hạn chế ghi tiến độ quá dày và cần cache nội dung khóa học.")
    add_bullets(doc, [
        "Tách read path: course bundle, lesson, block, manifest dùng snapshot/cache/CDN.",
        "Tách write path: progress, quiz, SCORM runtime ghi qua API/RPC có idempotency.",
        "Debounce hoặc batch progress 15 đến 30 giây/lần, ghi theo milestone thay vì mỗi tương tác nhỏ.",
        "Index bắt buộc: enrollment_id, profile_id, course_id, lesson_id, updated_at.",
        "Certificate, report và đồng bộ điểm chạy background job.",
    ])

    add_heading(doc, "5.2 VTraining", 2)
    add_body(doc, "VTraining có nhiều nghiệp vụ transaction-heavy như enrollment, attendance, assessment, certification. Query key bằng React Query đã rõ, nhưng màn class detail và dashboard có nguy cơ sinh nhiều query song song.")
    add_bullets(doc, [
        "Gom dữ liệu class detail bằng RPC/read model thay vì nhiều select rời từ client.",
        "Enrollment, attendance và assessment dùng transaction/RPC có idempotency.",
        "Dashboard dùng summary table hoặc materialized view.",
        "Báo cáo completion, đồng bộ quiz/reflection, gửi thông báo chuyển sang queue.",
        "Với lớp đông, phân trang roster và lazy-load kết quả theo tab.",
    ])

    add_heading(doc, "5.3 VSurvey", 2)
    add_body(doc, "VSurvey là workload ghi nhiều. Hiện public submission có thể insert trực tiếp vào bảng Supabase và export có luồng đọc toàn bộ submission theo page rồi tổng hợp phía client.")
    add_bullets(doc, [
        "Public submit đi qua API/RPC có rate limit, captcha hoặc honeypot cho link công khai.",
        "Thêm idempotency/unique guard theo form, respondent hoặc request id để chống submit trùng.",
        "Tách raw submissions và aggregated summary để dashboard không count raw table liên tục.",
        "Export lớn chạy async job và trả link file khi hoàn tất.",
        "Index: form_id + submitted_at desc, email, và GIN cho answers nếu query JSONB.",
    ])

    add_heading(doc, "5.4 VEvent", 2)
    add_body(doc, "VEvent đã đi đúng hướng khi dùng RPC cho public load, submit response, switch question và stats. Rủi ro lớn nhất là polling 5 giây và thời điểm spike khi cả phòng cùng submit.")
    add_bullets(doc, [
        "Submit vote cần unique constraint trên event_id, question_id, participant_key.",
        "Stats nên đọc từ summary table/counter, không group raw responses ở mỗi refresh.",
        "Audience dùng adaptive polling 10 đến 30 giây; presenter/monitor có thể dùng realtime hoặc polling nhanh hơn.",
        "Check-in hoặc slot giới hạn phải dùng transaction/RPC với lock để tránh race condition.",
        "Public join/submit cần Redis rate limit và chống spam theo IP, participant key, event code.",
    ])

    add_heading(doc, "6 Ước lượng hạ tầng cho mốc 3.000", 1)
    add_table(
        doc,
        ["Lớp hạ tầng", "Khuyến nghị ban đầu", "Ghi chú"],
        [
            ["Frontend/CDN", "Vercel/CDN, cache immutable cho asset, budget bundle theo route.", "Tĩnh là phần dễ scale nhất nếu media tách khỏi app."],
            ["API", "Serverless functions theo nhóm route, p95 mục tiêu dưới 500 đến 800 ms cho read path.", "Cần rate limit phân tán và timeout rõ."],
            ["Database", "Supabase Pro/Team trở lên, pooler bật, slow query dashboard.", "Giảm direct query bằng RPC/cache/read model."],
            ["Redis", "Upstash/Redis bắt buộc cho rate limit, counter, cache ngắn.", "Không dùng local memory cho production scale."],
            ["Queue/Worker", "Worker pool cho export, email, scoring, certificate, sync.", "Có retry, dead-letter và audit."],
            ["Storage/CDN", "Object storage cho video, SCORM, file; signed URL khi cần.", "Media không đi qua API app."],
        ],
        widths=[Inches(1.35), Inches(3.35), Inches(2.05)],
    )

    add_heading(doc, "7 Chiến lược kiểm thử tải", 1)
    add_table(
        doc,
        ["Loại test", "Mục tiêu", "Kịch bản đề xuất"],
        [
            ["Load test", "Xác nhận tải mục tiêu ổn định.", "3.000 users phân bổ theo VLearning, VTraining, VSurvey, VEvent; đo p95, p99, error rate."],
            ["Stress test", "Tìm điểm gãy.", "Tăng dần 500, 1.000, 2.000, 3.000, 4.000 users; theo dõi DB CPU, connection, PostgREST latency."],
            ["Spike test", "Mô phỏng cao điểm.", "VEvent mở câu hỏi và 1.000 đến 3.000 lượt submit trong 1 đến 3 phút."],
            ["Soak test", "Kiểm tra ổn định dài hạn.", "VLearning/VTraining chạy 2 đến 4 giờ với user xem bài, lưu tiến độ, làm quiz."],
            ["Failover test", "Kiểm tra resilience.", "Redis chậm/mất kết nối, Supabase timeout, queue retry, export lỗi."],
        ],
        widths=[Inches(1.25), Inches(2.3), Inches(3.25)],
    )

    add_heading(doc, "8 Lộ trình cải tiến", 1)
    add_table(
        doc,
        ["Giai đoạn", "Việc cần làm", "Kết quả mong đợi"],
        [
            ["Ngắn hạn", "Nâng target 3.000, bật Redis rate limit enforce, chạy k6 hot path, giảm polling, thêm index nóng.", "Có baseline capacity và giảm rủi ro nghẽn tức thời."],
            ["Trung hạn", "Chuẩn hóa BFF/API cho write path công khai, read model cho dashboard/report/stats, queue cho export/email/scoring.", "Hệ thống chịu tải tốt hơn, ít phụ thuộc client và giảm DB pressure."],
            ["Dài hạn", "Tách bounded context backend theo Learning, Training, Survey, Event; event-driven projection; analytics DB/read replica.", "Nền tảng mở rộng bền vững, dễ vận hành và bảo trì lâu dài."],
        ],
        widths=[Inches(1.2), Inches(3.55), Inches(2.0)],
    )

    add_heading(doc, "9 Checklist hành động", 1)
    add_bullets(doc, [
        "Xác định SLO production: p95, p99, error rate, login latency, submit latency.",
        "Bật Redis/Upstash distributed rate limit cho production.",
        "Chạy load test cho login, VLearning course load/progress, VSurvey submit, VEvent submit/stats.",
        "Xuất top slow queries từ Supabase và bổ sung index.",
        "Chuyển export/report/email/scoring sang queue.",
        "Tạo read model cho survey summary, event stats, class completion và dashboard.",
        "Thiết lập Sentry/APM, log correlation, Supabase query insights và alert.",
        "Viết runbook rollback, backup, restore drill và incident response.",
    ])

    add_heading(doc, "10 Thông tin cần làm rõ", 1)
    add_bullets(doc, [
        "Số concurrent users thực tế theo từng module và phân bố hành vi người dùng.",
        "Supabase plan hiện tại, giới hạn connection, PostgREST quota và database size.",
        "Kết quả load test gần nhất, nếu đã chạy.",
        "Dung lượng video, SCORM, file upload và chính sách CDN hiện tại.",
        "Danh sách bảng production lớn nhất, index hiện có và slow query logs.",
        "Vercel function metrics: invocation count, duration, cold start, error rate.",
        "Yêu cầu SLA/RTO/RPO và chính sách backup/restore.",
    ])

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Vinabrain Architecture Review | 3.000 concurrent users")
    set_run_font(footer_run, size=8.5, color="666666")

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_report()
    print(output)

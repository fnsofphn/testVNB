"""Real-corpus regressions plus server workflow/permission tests in an isolated DB."""
import os, sys, tempfile, unittest, copy, json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'vcoaching'))
TEMP = tempfile.TemporaryDirectory(prefix='vcoaching-test-')
os.environ['VCOACHING_DATA_DIR'] = TEMP.name
import server as s
from documents import parse, STEPS, norm, business_code, export_docx, mapped_outputs, numeric_conflicts, rules
CORPUS = Path(r'D:\03. Data\2026\V-Coaching\Tài liệu\2. WS1A_CHI TIẾT THEO PHIÊN')
MASTER = Path(r'D:\04. Code\Vcoaching\CLSP_Mau_01_Master_CTHD_VNPT_Rising_2026_2028.docx')

def who(role='data', super_admin=False, unit='vcoaching-test-unit'):
 return {'id':role,'email':role+'@test.invalid','role':role,'super':super_admin,'switched':False,'projects':['vcoaching-test'],'units':[unit],'initiatives':[],'assignment_actor':role}

class Workflow(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  s.ensure_seed(); cls.parsed=parse(MASTER); cls.client=s.app.test_client()
 def setUp(self):
  s.DB.execute("DELETE FROM records WHERE kind NOT IN ('project','unit')")
  f={'id':'source','type':'file','project':'vcoaching-test','unit':'vcoaching-test-unit','name':MASTER.name,'sha256':self.parsed['sha256'],'status':'parsed'}
  s.put('file',f); self.iid=s.add_candidate(copy.deepcopy(self.parsed['candidates'][0]),f)
 def request(self,op,role='data',status=200,**body):
  with patch.object(s,'actor',return_value=who(role)):
   r=self.client.post('/vc-api/'+op,json={'id':self.iid,**body})
  self.assertEqual(r.status_code,status,r.get_json()); return r.get_json()
 def test_workflow_and_private_review(self):
  self.request('edit',reason='Đã đối chiếu nguồn CLSP',resolve_issues=True)
  self.request('confirm',reason='Đã đối chiếu đủ')
  self.request('assign',role='unit',status=403,experts=['unit'])
  item=s.get(self.iid); item['experts']=['expert']; s.put('initiative',item)
  self.request('comment',role='expert',step='7A',problem='Cần xác nhận',why='Căn cứ nguồn chưa rõ',question='Ai làm khác?',notes='PRIVATE',internal=True)
  with patch.object(s,'actor',return_value=who('unit')):
   visible=self.client.get('/vc-api/detail?id='+self.iid).get_json()['initiative']
  self.assertEqual(visible['comments'],[])
  for c in s.get(self.iid)['comments']:
   if c.get('stale'):continue
   self.request('review',role='expert',comment_id=c['id'],state='approved',internal=False)
  self.request('lock',role='expert'); self.request('release',role='project')
  self.request('comment',role='expert',status=400,step='7A',problem='x',why='x',question='x')
  self.request('respond',role='unit',status=400,agreement='partial',reason='')
  self.request('respond',role='unit',agreement='agree',revisions={'7A':'Đơn vị tự hiệu chỉnh'},self_rating={'gate':{'1':'Đạt'}},submit=True)
  self.request('respond',role='unit',status=400,agreement='agree',submit=True)
  self.request('release',role='project',status=400)
  self.request('confirm',reason='Không được bỏ qua bản nộp',status=400)
  self.request('unlock',role='expert',reason='Không được mở bản nộp',status=400)
  self.request('recheck',role='data')
  self.assertEqual(s.get(self.iid)['locked_steps'],[])
  self.request('finalize',reason='Chưa duyệt nhận xét mới',status=400)
  for c in s.get(self.iid)['comments']:
   if not c.get('stale'): self.request('review',role='expert',comment_id=c['id'],state='approved')
  self.request('confirm',reason='Xác nhận sau kiểm tra lại')
  self.assertEqual(s.get(self.iid)['status'],'rechecked')
  self.request('finalize',reason='Đã duyệt kiểm tra lại')
  self.assertEqual(s.get(self.iid)['status'],'complete')
  self.request('unlock',role='system',target='unit',reason='Cho đơn vị bổ sung nguồn')
  self.assertEqual(s.get(self.iid)['status'],'released')
 def test_partial_step_release_and_private_fields(self):
  self.request('edit',reason='Đối chiếu',resolve_issues=True)
  self.request('confirm',reason='Đủ căn cứ')
  item=s.get(self.iid); item['experts']=['expert']; item['comments']=[]; s.put('initiative',item)
  for step in ('1','2'):
   self.request('comment',role='expert',step=step,problem='Nhận định',why='Nguồn',question='Câu hỏi?',hypothesis='PRIVATE',internal=False)
  for c in s.get(self.iid)['comments']:
   self.request('review',role='expert',comment_id=c['id'],state='approved',internal=False)
  self.request('lock',role='expert',step='1')
  self.request('release',role='system',step='2',status=400)
  self.request('release',role='system',step='1')
  visible=s.visible_initiative(who('unit'),s.get(self.iid))
  self.assertEqual([c['step'] for c in visible['comments']],['1'])
  self.assertNotIn('hypothesis',visible['comments'][0])
  self.request('respond',role='unit',agreement='agree',revisions={'2':'Không được sửa'},status=403)
  self.request('comment',role='expert',step='2',problem='Còn mở',why='Nguồn',question='Câu hỏi?')
  self.request('comment',role='expert',step='1',problem='Đã khóa',why='Nguồn',question='?',status=400)
  self.request('unlock',role='expert',step='1',reason='Kiểm tra lại')
  self.assertEqual(s.visible_initiative(who('unit'),s.get(self.iid))['comments'],[])
 def test_submission_locks_only_selected_step(self):
  self.request('edit',reason='Đối chiếu',resolve_issues=True)
  self.request('confirm',reason='Đủ căn cứ')
  item=s.get(self.iid); item['experts']=['expert']; item['comments']=[]; s.put('initiative',item)
  for step in ('1','2'):
   self.request('comment',role='expert',step=step,problem='Nhận định',why='Nguồn',question='Câu hỏi?',internal=False)
  for c in s.get(self.iid)['comments']:
   self.request('review',role='expert',comment_id=c['id'],state='approved',internal=False)
  for step in ('1','2'):
   self.request('lock',role='expert',step=step)
   self.request('release',role='project',step=step)
  self.request('respond',role='unit',step='1',agreement='agree',revisions={'1':'Bản sửa 1'},submit=True)
  self.assertEqual(s.get(self.iid)['status'],'self_review')
  self.assertEqual(s.get(self.iid)['submitted_steps'],['1'])
  self.request('release',role='project',step='1',status=400)
  current=s.get(self.iid)
  self.assertTrue(all(c['version']==current['versions'][-1]['number'] for c in current['comments'] if not c.get('stale')))
  self.assertEqual(len(s.visible_initiative(who('unit'),s.get(self.iid))['comments']),2)
  self.request('respond',role='unit',step='1',agreement='agree',revisions={'1':'Nộp lại'},submit=True,status=403)
  self.request('respond',role='unit',step='2',agreement='agree',revisions={'2':'Bản sửa 2'},submit=True)
  self.assertEqual(s.get(self.iid)['status'],'submitted')
  self.assertEqual(s.get(self.iid)['versions'][-1]['revisions'],{'1':'Bản sửa 1','2':'Bản sửa 2'})
 def test_bottleneck_self_assessment_is_saved_with_step_four(self):
  self.request('edit',reason='Đối chiếu',resolve_issues=True)
  self.request('confirm',reason='Đủ căn cứ')
  item=s.get(self.iid); item['experts']=['expert']; item['comments']=[]; s.put('initiative',item)
  self.request('comment',role='expert',step='4',problem='Điểm nghẽn cần soi',why='Căn cứ nguồn',question='Điều gì lặp lại?',internal=False)
  comment=s.get(self.iid)['comments'][-1]
  self.request('review',role='expert',comment_id=comment['id'],state='approved',internal=False)
  self.request('lock',role='expert',step='4')
  self.request('release',role='project',step='4')
  revisions={'4':'Đơn vị bổ sung điểm nghẽn'}
  self.request('respond',role='unit',step='4',agreement='agree',revisions=revisions,bottleneck_check={'1':'Có','2':'Chưa rõ'},submit=False)
  self.assertEqual(s.get(self.iid)['responses'][-1]['bottleneck_check'],{'1':'Có','2':'Chưa rõ'})
  self.request('respond',role='expert',step='4',agreement='agree',revisions=revisions,bottleneck_check={'1':'Có'},status=403)
  self.request('respond',role='unit',step='4',agreement='agree',revisions=revisions,bottleneck_check={'6':'Có'},status=400)
  self.request('respond',role='unit',step='4',agreement='agree',revisions=revisions,bottleneck_check={'1':'Có','2':'Không','3':'Có','4':'Có','5':'Có'},submit=True)
  self.assertEqual(s.get(self.iid)['versions'][-1]['bottleneck_check']['2'],'Không')
  self.request('respond',role='unit',step='4',agreement='agree',revisions=revisions,bottleneck_check={'1':'Có'},status=400)
 def test_extraction_correction_preserves_original_and_scope(self):
  r=s.get(self.iid); b=r['versions'][-1]['blocks'][0]; original=b['text']
  self.request('edit',reason='Sửa lỗi trích xuất theo nguồn',corrections={b['id']:'Đã đối chiếu'})
  after=s.get(self.iid); changed=after['versions'][-1]['blocks'][0]
  self.assertEqual(changed['original_text'],original)
  self.assertEqual(after['versions'][0]['blocks'][0]['text'],original)
  self.request('edit',reason='Sai phạm vi',corrections={'foreign':'bad'},status=400)
  self.request('edit',role='unit',reason='Không đúng quyền',corrections={b['id']:'bad'},status=403)
 def test_config_and_comparison_permissions(self):
  self.request('config',role='unit',project='vcoaching-test',stage='Trích xuất',text='Giữ nguồn',status=403)
  self.request('config',role='system',project='vcoaching-test',stage='Trích xuất',text='Giữ nguồn')
  self.assertEqual(s.rows('config')[0]['text'],'Giữ nguồn')
  self.request('compare-review',role='unit',step='1',label='Rõ hơn',reason='Bằng chứng',status=403)
  self.request('compare-review',role='project',step='1',label='Rõ hơn',reason='Bằng chứng')
  self.assertEqual(s.get(self.iid)['comparison']['1']['version'],1)
 def test_hr_multiple_files_keep_three_initiatives_and_extra_proposal(self):
  s.DB.execute("DELETE FROM records WHERE kind NOT IN ('project','unit')")
  directory=next(p for p in CORPUS.rglob('1. Ban Nhân lực') if p.is_dir())
  expected={}
  for path in sorted(directory.glob('*.docx')):
   parsed=parse(path); f={'id':path.stem,'type':'file','project':'vcoaching-test','unit':'vcoaching-test-unit','name':path.name,'sha256':parsed['sha256'],'status':'parsed'}
   s.put('file',f)
   for candidate in parsed['candidates']:
    rid=s.add_candidate(candidate,f); expected[rid]={b['id'] for b in candidate['blocks']}
  records=s.rows('initiative')
  masters=[r for r in records if 'master' in r['forms']]
  details=[r for r in records if 'detail' in r['forms']]
  self.assertEqual(len(masters),3);self.assertEqual(len(details),4)
  before=copy.deepcopy(records)
  automatic=mapped_outputs(records)
  self.assertEqual(records,before)
  self.assertEqual(len(automatic),4)
  self.assertEqual(sorted(len(r['files']) for r in automatic),[1,2,2,2])
  self.assertTrue(any('Khác biệt nhận diện' in n for r in automatic for n in r['mapping_notes']))
  with patch.object(s,'actor',return_value=who('unit')):
   output=self.client.get('/vc-api/mapped-export?files='+','.join(f['id'] for f in s.rows('file')))
  self.assertEqual(output.status_code,200)
  import io
  with ZipFile(io.BytesIO(output.data)) as z:
   self.assertEqual(len([n for n in z.namelist() if n.endswith('.docx')]),4)
   self.assertIn('AI-Ready',z.read('Tong-hop-ket-qua.txt').decode())
  with patch.object(s,'actor',return_value=who('unit',unit='other')):
   denied=self.client.get('/vc-api/mapped-export?files='+s.rows('file')[0]['id'])
  self.assertEqual(denied.status_code,403)
  for master in masters:
   anchor=master['name'].split(' – ')[0]
   detail=next(r for r in details if r['name'].startswith(anchor))
   self.iid=master['id']
   self.request('convert',files=[f['id'] for f in s.rows('file')])
   merged=s.get(self.iid)
   self.assertEqual(set(merged['forms']),{'master','detail'})
   self.assertEqual(len(merged['files']),2)
   self.assertEqual(merged['code'],master['code'])
   self.assertEqual({b['id'] for b in merged['versions'][-1]['blocks']},expected[master['id']]|expected[detail['id']])
   self.assertTrue(all(merged['versions'][-1]['fields'][k] for k in ('1','2','3','4','6','7A','7B','7C','8')))
  active=[r for r in s.rows('initiative') if not r.get('merged_into')]
  self.assertEqual(len(active),4)
  proposal=next(r for r in active if r['name'].startswith('AI-Ready'))
  self.assertEqual(len(proposal['files']),1)
  self.assertIn('missing_code',proposal['issues'])
  versions={r['id']:len(r['versions']) for r in active}
  self.request('convert',files=[f['id'] for f in s.rows('file')])
  self.assertEqual(versions,{r['id']:len(r['versions']) for r in s.rows('initiative') if not r.get('merged_into')})
  with patch.object(s,'actor',return_value=who('data')):
   workspace=self.client.get('/vc-api/workspace').get_json()
  self.assertEqual(len(workspace['initiatives']),4)
 def test_new_upload_does_not_overwrite_submitted_record(self):
  original=s.get(self.iid); original.update(status='submitted',released=True)
  original['versions'][-1]['revisions']={'1':'Nội dung đơn vị đã nộp'}
  s.put('initiative',original)
  candidate=copy.deepcopy(self.parsed['candidates'][0]); candidate.update(id='new-upload-record',issues=[])
  file={**s.get('source','file'),'id':'second-file'}; s.put('file',file)
  result=s.add_candidate(candidate,file)
  self.assertNotEqual(result,self.iid)
  self.assertEqual(s.get(self.iid),original)
  self.assertIn('code_conflict',s.get(result)['issues'])
 def test_numeric_source_conflicts_preserve_evidence(self):
  r=s.get(self.iid)
  blocks=[{'id':'a','file_id':'one','label':'Baseline','text':'5.2%'}, {'id':'b','file_id':'two','label':'Baseline','text':'52%'}]
  r['versions'][-1]['blocks']=blocks; r['versions'][-1]['fields']={'3':['a','b']}
  before=copy.deepcopy(r)
  conflicts=numeric_conflicts(r)
  self.assertEqual(len(conflicts),1); self.assertEqual(r,before)
  alerts=[c for c in rules(r) if c['rule']=='NUMERIC_SOURCE_DIFFERENCE_V1']
  self.assertEqual(alerts[0]['evidence'],blocks); self.assertEqual(alerts[0]['state'],'pending')
  blocks[1]['text']='5.2%'; self.assertEqual(numeric_conflicts(r),[])
  blocks[1].update(text='52%',file_id='one'); self.assertEqual(numeric_conflicts(r),[])
 def test_worker_progress_waits_for_other_received_files(self):
  f=s.get('source','file'); f.update(status='reading',stored='worker-test.docx')
  (s.DATA/'sources'/f['stored']).write_bytes(MASTER.read_bytes()); s.put('file',f)
  waiting={**f,'id':'waiting','status':'queued'}; s.put('file',waiting)
  s.process_file(f)
  done=s.get('source','file')
  self.assertEqual(done['status'],'parsed')
  self.assertEqual(done['progress']['extract']['status'],'done')
  self.assertEqual(done['progress']['merge']['status'],'awaiting_selection')
  self.request('convert',files=['source','waiting'],status=409)
  self.request('convert',files=['source'])
  waiting['status']='error'; s.put('file',waiting)
  self.request('convert',files=['source'])
  self.assertEqual(s.get('source','file')['progress']['result']['status'],'done')
  broken={**f,'id':'broken','stored':'missing.docx','status':'reading'}; s.put('file',broken)
  s.process_file(broken)
  self.assertEqual(s.get('broken','file')['status'],'error')
  self.assertEqual(s.get('broken','file')['progress']['extract']['status'],'error')
 def test_scope_and_forged_mapping(self):
  item=s.get(self.iid)
  self.assertFalse(s.can(who('unit',unit='other'),item))
  self.assertFalse(s.can(who('expert'),item))
  self.assertTrue(s.can(who('project'),item))
  fields={k:[] for k in STEPS}; fields['1']=['foreign-block']
  self.request('edit',status=400,fields=fields,reason='bad')
  self.request('account',role='system',status=403)
 def test_automatic_mapping_separates_units_and_ambiguous_names(self):
  master=s.get(self.iid)
  detail=copy.deepcopy(master); detail.update(id='detail',forms=['detail'],files=['detail-file'])
  other=copy.deepcopy(master); other.update(id='other-unit',unit='other-unit',files=['other-file'])
  results=mapped_outputs([master,detail,other])
  self.assertEqual(len(results),2)
  self.assertEqual(next(r for r in results if r['unit']=='other-unit')['files'],['other-file'])
  self.assertEqual(set(next(r for r in results if r['unit']==master['unit'])['files']),set(master['files']+['detail-file']))
  duplicate=copy.deepcopy(master); duplicate['id']='ambiguous-master'
  self.assertEqual(len(mapped_outputs([master,duplicate,detail])),1)
  f=s.get('source','file'); f['status']='reading'; s.put('file',f)
  with patch.object(s,'actor',return_value=who('unit')):
   response=self.client.get('/vc-api/mapped-export?files=source')
  self.assertEqual(response.status_code,409)
 def test_selected_files_only_and_repeat_conversion(self):
  original=s.get(self.iid)
  selected=copy.deepcopy(original); selected.update(id='selected',files=['selected-file'],forms=['detail'])
  unselected=copy.deepcopy(original); unselected.update(id='unselected',files=['unselected-file'],forms=['detail'])
  for r in (selected,unselected):
   for block in r['versions'][-1]['blocks']: block['id']=r['id']+block['id']
   s.put('initiative',r)
   s.put('file',{**s.get('source','file'),'id':r['files'][0]})
  before=s.get('unselected')
  self.request('convert',files=['source','selected-file'])
  active=[r for r in s.rows('initiative') if not r.get('merged_into') and not r.get('conversion_pending')]
  self.assertEqual(len(active),1)
  self.assertEqual(set(active[0]['files']),{'source','selected-file'})
  self.assertEqual(s.get('unselected'),before)
  version=active[0]['versions'][-1]['number']
  self.request('convert',files=['source','selected-file'])
  self.assertEqual(s.get(active[0]['id'])['versions'][-1]['number'],version)
  self.request('convert',files=['source'])
 def test_abbreviations_typos_and_detail_only_grouping(self):
  base=s.get(self.iid); base.update(forms=['detail'],name='Nâng cao chất lượng chăm sóc khách hàng doanh nghiệp')
  alias=copy.deepcopy(base); alias.update(id='alias',files=['alias-file'],name='Nâng cao chất lượng CSKH DN')
  typo=copy.deepcopy(base); typo.update(id='typo',files=['typo-file'],name='Nâng cao chất lương chăm sóc khách hàng doanh nghiệp')
  self.assertEqual(len(mapped_outputs([base,alias,typo])),1)
  different=copy.deepcopy(base); different.update(id='different',name='Nâng cao chất lượng đào tạo nhân lực')
  self.assertEqual(len(mapped_outputs([base,different])),2)
  other=copy.deepcopy(alias); other.update(id='other',unit='another-unit')
  self.assertEqual(len(mapped_outputs([base,other])),2)
  numbered=copy.deepcopy(base); numbered.update(id='numbered',name=base['name']+' 2')
  self.assertEqual(len(mapped_outputs([base,numbered])),2)
 def test_fanpage_short_title_real_sources(self):
  directory=next(p for p in CORPUS.rglob('4. Ban Truyền thông') if p.is_dir())
  records=[]
  for filename in ('Mau_01_Master_CTHD_VNPT_Rising_2026_2028_Ban TT.docx','SK02_Mau_02_CTHD_VNPT_Rising_2026_2028_Ban TT.docx'):
   parsed=parse(directory/filename)
   file={**s.get('source','file'),'id':filename,'name':filename,'sha256':parsed['sha256']}
   s.put('file',file)
   for c in parsed['candidates']:
    if business_code(c['code'])=='SK02': records.append(s.get(s.add_candidate(c,file)))
  self.assertEqual(len(records),2)
  original=copy.deepcopy(records)
  result=mapped_outputs(records)
  self.assertEqual(len(result),1)
  self.assertEqual(result[0]['name'],records[0]['name'])
  self.assertEqual(len(result[0]['files']),2)
  self.assertEqual(len(result[0]['versions'][-1]['blocks']),sum(len(r['versions'][-1]['blocks']) for r in records))
  self.assertEqual(records,original)
  self.request('convert',files=[r['files'][0] for r in records])
  self.assertEqual(len([r for r in s.rows('initiative') if not r.get('merged_into') and r['id'] in {v['id'] for v in records}]),1)
  wrong=copy.deepcopy(original[1]); wrong['code']='SK03'
  self.assertEqual(len(mapped_outputs([original[0],wrong])),2)
  wrong=copy.deepcopy(original[1]); wrong['source_unit_name']='Ban Nhân lực'
  self.assertEqual(len(mapped_outputs([original[0],wrong])),2)
  wrong=copy.deepcopy(original[1]); wrong['versions'][-1]['blocks']=[]
  self.assertEqual(len(mapped_outputs([original[0],wrong])),2)
  second=copy.deepcopy(original[0]); second.update(id='another-master',name=original[0]['name']+', phạm vi khác')
  result=mapped_outputs([original[0],second,original[1]])
  # If both full titles themselves match, they are one source group; otherwise
  # the short title must not choose arbitrarily between distinct groups.
  self.assertTrue(all(len(r['contributing_ids'])!=2 or original[1]['id'] not in r['contributing_ids'] for r in result))
 def test_catalog_and_legacy_reporting_preserve_authorization(self):
  units=s.rows('unit'); names={u['name'] for u in units}
  self.assertTrue({'Ban Nhân lực','Ban Truyền thông','VNPT Hà Nội'}<=names)
  communication=next(u for u in units if u['name']=='Ban Truyền thông')
  self.assertEqual(s.get(communication['id'],'unit')['name'],'Ban Truyền thông')
  r=s.get(self.iid); r['unit_name']='Ban Truyền thông'
  before=copy.deepcopy(r)
  visible=s.visible_initiative(who('data'),r)
  self.assertEqual(visible['report_unit'],communication['id'])
  self.assertEqual(visible['unit'],r['unit'])
  self.assertEqual(visible['unit_type'],'Ban chức năng')
  self.assertEqual(r,before)
  self.assertFalse(s.can(who('unit'),communication))
  self.assertTrue(s.can(who('project'),communication))
  # Catalog metadata never grants a data/unit account access to a new unit.
  with s.app.test_request_context():
   with self.assertRaises(s.Problem):s.scoped(who('unit'),communication['id'],'unit')
  s.put('unit',{**communication,'name':'Tên quản trị đã sửa'})
  self.assertEqual(len([u for u in s.rows('unit') if u['id']==communication['id']]),1)
  self.assertEqual(s.get(communication['id'])['name'],'Tên quản trị đã sửa')
  s.DB.execute('DELETE FROM records WHERE id=?',(communication['id'],))
 def test_partial_master_selection_and_stable_tracking_numbers(self):
  base=s.get(self.iid)
  s.DB.execute('DELETE FROM records WHERE id=?',(self.iid,))
  first={**copy.deepcopy(base),'id':'first','files':['source','detail-one'],'forms':['master','detail'],'conversion_pending':False}
  second={**copy.deepcopy(base),'id':'second','files':['source','detail-two'],'forms':['master','detail'],'conversion_pending':False,'unit_name':'Ban khác'}
  child={**copy.deepcopy(base),'id':'child','files':['detail-one'],'forms':['detail'],'merged_into':'first'}
  for r in (first,second,child):s.put('initiative',r)
  for fid in ('detail-one','detail-two'):s.put('file',{**s.get('source','file'),'id':fid})
  result=self.request('convert',files=['source','detail-one'])
  self.assertEqual(result['initiatives'],['first'])
  self.assertEqual(s.get('second')['versions'],second['versions'])
  codes=[s.get(key)['tracking_code'] for key in ('first','second')]
  self.assertEqual(len(set(codes)),2)
  self.request('convert',files=['source','detail-one'])
  self.assertEqual([s.get(key)['tracking_code'] for key in ('first','second')],codes)
  self.assertEqual(s.get('first')['code'],base['code'])
  third={**copy.deepcopy(base),'id':'third','files':['later-file'],'conversion_pending':True}
  s.put('initiative',third);s.put('file',{**s.get('source','file'),'id':'later-file'})
  self.request('convert',files=['later-file'])
  self.assertNotIn(s.get('third')['tracking_code'],codes)
  self.assertGreater(int(s.get('third')['tracking_code'][3:]),max(int(code[3:]) for code in codes))
 def test_fresh_grant_and_switch_rejects_forgery(self):
  user={'id':'admin','email':'admin@test.invalid','app_metadata':{'vcoaching':{'active':True,'role':'system','super_admin':True,'projects':[],'units':[]}}}
  def auth(path,**kwargs):
   return [{'id':'p','active':True}] if path.startswith('/rest/') else copy.deepcopy(user)
  with s.app.test_request_context(headers={'Authorization':'Bearer test','X-VCoaching-Role':'unit','X-VCoaching-Project':'vcoaching-test','X-VCoaching-Unit':'vcoaching-test-unit'}),patch.object(s,'sb',side_effect=auth):
   a=s.actor(); self.assertFalse(a['super']); self.assertTrue(a['switched']); self.assertEqual(a['units'],['vcoaching-test-unit'])
   user['app_metadata']['vcoaching']['super_admin']=False
   with self.assertRaises(s.Problem):s.actor()
   user['app_metadata']['vcoaching']['active']=False
   with self.assertRaises(s.Problem):s.actor()
 def test_export_preserves_package_and_source(self):
  r=s.get(self.iid); out=Path(TEMP.name)/'WS1b.docx'; export_docx(r,s.TEMPLATE,out)
  from docx import Document
  d=Document(out); self.assertEqual(len(d.tables),4); self.assertEqual(len(d.tables[1].rows),11)
  text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
  self.assertNotIn('TP.HCM',text); self.assertIn('PHIẾU NHÁP',text)
  self.assertIn('7A. HÀNH VI MỚI',text); self.assertIn('7B. CƠ CHẾ MỚI',text)
  self.assertNotIn('[x]', '\n'.join(c.text for t in d.tables[2:] for row in t.rows for c in row.cells))
  for block in r['versions'][-1]['blocks']:
   if block['text'].strip():self.assertIn(block['text'],text)
  with ZipFile(s.TEMPLATE) as a,ZipFile(out) as b:
   self.assertEqual(set(a.namelist()),set(b.namelist()))
   for name in a.namelist():
    if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name),name)
 def test_last_admin_protected(self):
  admin={'id':'admin','email':'admin@test.invalid','app_metadata':{'vcoaching':{'role':'system','active':True,'super_admin':True,'projects':['vcoaching-test'],'units':['vcoaching-test-unit']}}}
  with s.app.test_request_context(),patch.object(s,'sb',return_value=admin),patch.object(s,'account_users',return_value=[admin]):
   with self.assertRaisesRegex(s.Problem,'Admin tổng cuối cùng'):
    s.accounts('account',who('system',True),{'user_id':'admin','action':'lock','confirmed':True,'reason':'test'})

class Corpus(unittest.TestCase):
 def setUp(self): s.ensure_seed()
 def test_real_forms(self):
  master=parse(MASTER); self.assertEqual([c['code'] for c in master['candidates']],['SK2.1','SK2.2','SK2.3'])
  cases=[('Mau_02_Sang_kien_PCTT_V2.pdf',3,3),('PL_Du thao CTHD 2026-2028 VNPT Rising_Ban PTTT.pdf',10,5)]
  for filename,forms,unique in cases:
   p=next(CORPUS.rglob(filename)); result=parse(p); self.assertEqual(len(result['candidates']),forms)
   s.DB.execute("DELETE FROM records WHERE kind='initiative'")
   f={'id':filename,'project':'vcoaching-test','unit':'vcoaching-test-unit','name':filename,'sha256':result['sha256']}
   for c in result['candidates']:s.add_candidate(c,f)
   self.assertEqual(len(mapped_outputs(s.rows('initiative'))),unique,filename)
  for p in CORPUS.rglob('*.docx'):
   if 'SK05' in p.name:
    result=parse(p)
    if any(business_code(c['code'])=='SK01' for c in result['candidates']):
     self.assertTrue(any('filename_code_conflict' in c['issues'] for c in result['candidates']))
  scan=next(CORPUS.rglob('CNNV CL.PDF'),None)
  if scan:self.assertEqual(parse(scan)['status'],'needs_ocr')

if __name__=='__main__':
 try: unittest.main(verbosity=2)
 finally: s.DB.close()

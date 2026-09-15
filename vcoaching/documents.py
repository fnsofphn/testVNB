"""Local, deterministic extraction. Source text is data, never instructions."""
import hashlib
import re
import unicodedata
import uuid
from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
from functools import lru_cache
from lxml import etree
from docx import Document
import pymupdf

STEPS = {'1': 'Mục tiêu Rising và vai trò', '2': 'Trạng thái đích',
         '3': 'Hiện trạng và bằng chứng', '4': 'Điểm nghẽn thật',
         '5': 'Đối tượng bị ảnh hưởng', '6': 'Sáng kiến can thiệp',
         '7A': 'Hành vi mới', '7B': 'Cơ chế mới', '7C': 'Kết quả mới',
         '8': 'Thực thi và bằng chứng'}


def norm(text):
    text = unicodedata.normalize('NFD', str(text or '').lower()).replace('đ', 'd')
    return re.sub(r'\s+', ' ', ''.join(c for c in text if unicodedata.category(c) != 'Mn')).strip()


def blank(text):
    value = re.sub(r'[.\s…_\-\[\]☐:]+', '', text or '')
    return not value or norm(text).startswith(('dien noi dung tai day', 'ky, ghi ro ho ten'))


def business_code(text):
    match = re.search(r'(?i)SK\s*\d+(?:[.,]\d+)?', text or '')
    return re.sub(r'\s', '', match[0]).upper().replace(',', '.') if match else ''


def block(text, locator, label='', kind='content'):
    return {'id': str(uuid.uuid4()), 'text': text or '', 'label': label,
            'locator': locator, 'kind': kind}


def read_docx(path):
    with ZipFile(path) as z:
        if sum(x.file_size for x in z.infolist()) > 120 * 1024 * 1024:
            raise ValueError('DOCX giải nén vượt giới hạn 120 MB')
    doc = Document(path)
    tables, paragraphs, order = [], [], []
    ti = pi = 0
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    section = ''
    for element in doc.element.body:
        if element.tag.endswith('}p'):
            pi += 1
            text = Paragraph(element, doc).text
            if re.match(r'^[IVX]+[.\s]', text.strip()):
                section = text
            b = block(text, {'format': 'docx', 'paragraph': pi, 'section': section})
            paragraphs.append(b)
            order.append(('paragraph', b))
        elif element.tag.endswith('}tbl'):
            ti += 1
            rows = []
            for ri, row in enumerate(Table(element, doc).rows, 1):
                cells, seen = [], set()
                for ci, cell in enumerate(row.cells, 1):
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    cells.append(block(cell.text, {'format': 'docx', 'section': section,
                        'table': ti, 'row': ri, 'column': ci,
                        'paragraphs': list(range(1, len(cell.paragraphs) + 1))}))
                rows.append(cells)
            t = {'rows': rows, 'index': ti}
            tables.append(t)
            order.append(('table', t))
    return {'tables': tables, 'paragraphs': paragraphs, 'order': order, 'ocr_pages': [], 'pages': None}


def read_pdf(path):
    doc = pymupdf.open(path)
    if doc.is_encrypted:
        raise ValueError('PDF có mật khẩu; cần bản có thể đọc')
    tables, paragraphs, order, ocr = [], [], [], []
    for pn, page in enumerate(doc, 1):
        lines = []
        for b in page.get_text('dict')['blocks']:
            for line in b.get('lines', []):
                # Rotated watermark is retained separately, never mixed into answers.
                text = ''.join(s['text'] for s in line['spans'])
                lines.append((text, line['bbox'], line['dir']))
        horizontal = [x for x in lines if abs(x[2][1]) < .05]
        if sum(len(x[0].strip()) for x in horizontal) < 30:
            ocr.append(pn)
        table_items = []
        for ti, t in enumerate(page.find_tables().tables, 1):
            rows = []
            for ri, row in enumerate(t.rows, 1):
                cells = []
                for ci, rect in enumerate(row.cells, 1):
                    if rect is None:
                        continue
                    x0, y0, x1, y1 = rect
                    text = '\n'.join(s for s, bb, _ in horizontal
                        if x0 - 1 <= (bb[0] + bb[2]) / 2 <= x1 + 1
                        and y0 - 1 <= (bb[1] + bb[3]) / 2 <= y1 + 1)
                    cells.append(block(text, {'format': 'pdf', 'page': pn, 'table': ti,
                        'row': ri, 'column': ci, 'bbox': list(rect)}))
                if cells:
                    rows.append(cells)
            item = {'rows': rows, 'index': len(tables) + 1}
            tables.append(item)
            table_items.append((t.bbox[1], 'table', item, t.bbox))
        items = [(y, kind, t) for y, kind, t, _ in table_items]
        for text, bb, direction in lines:
            inside = any(rect[0] - 1 <= (bb[0] + bb[2]) / 2 <= rect[2] + 1
                         and rect[1] - 1 <= (bb[1] + bb[3]) / 2 <= rect[3] + 1
                         for _, _, _, rect in table_items)
            if inside and abs(direction[1]) < .05:
                continue
            b = block(text, {'format': 'pdf', 'page': pn, 'bbox': list(bb)},
                      kind='watermark' if abs(direction[1]) >= .05 else 'content')
            paragraphs.append(b)
            items.append((bb[1], 'paragraph', b))
        order.extend((kind, t) for _, kind, t in sorted(items, key=lambda x: x[0]))
    return {'tables': tables, 'paragraphs': paragraphs, 'order': order, 'ocr_pages': ocr, 'pages': len(doc)}


def new_candidate(form, code='', name=''):
    return {'id': str(uuid.uuid4()), 'form': form, 'code': code, 'name': name, 'unit_name': '',
            'blocks': [], 'fields': {s: [] for s in STEPS}, 'meta': {}, 'issues': []}


@lru_cache(maxsize=4)
def assessment_schema(template):
    doc = Document(template)
    return {group: [{'id': str(i + 1), 'question': row.cells[1].text}
                    for i, row in enumerate(doc.tables[index].rows[start:])]
            for group, index, start in [('gate', 2, 0), ('criteria', 3, 1)]}


def add(candidate, b, label, target=None, meta=None):
    b = {**b, 'label': label}
    candidate['blocks'].append(b)
    if not blank(b['text']):
        if target:
            candidate['fields'][target].append(b['id'])
        if meta:
            candidate['meta'].setdefault(meta, []).append(b['id'])


def label_target(label):
    n = norm(label)
    if n.startswith(('cau phan vnpt rising lien ket chinh', 'lien ket chinh', 'nguon va lien ket')): return '1'
    if n.startswith(('muc tieu hoac trang thai', 'muc tieu/trang thai', 'trang thai/thay doi')): return '2'
    if n.startswith(('baseline va nguon', 'van de/co hoi va hien trang', 'hien trang can thay doi')): return '3'
    if n.startswith('van de/diem nghen'): return '4'
    if n.startswith('doi tuong bi anh huong'): return '5'
    if n.startswith(('ten sang kien', 'sang kien/nhiem vu trong tam', 'hanh dong cu the')): return '6'
    if n.startswith(('ket qua cuoi cung', 'kpi, tieu chi va bang chung chinh')): return '7C'
    if any(n.startswith(x) for x in ['owner', 'lanh dao phu trach/bao tro', 'cac don vi phoi hop', 'don vi phoi hop', 'thoi gian thuc hien', 'cac moc ra soat', 'lo trinh', 'nguon luc']): return '8'
    return None


def parse(path, filename=None):
    path = Path(path)
    raw = read_docx(path) if path.suffix.lower() == '.docx' else read_pdf(path)
    all_blocks = raw['paragraphs'] + [b for t in raw['tables'] for r in t['rows'] for b in r]
    full = '\n'.join(b['text'] for b in all_blocks)
    candidates, current, unit_name = [], None, ''
    for t in raw['tables']:
        for row in t['rows']:
            if len(row) == 2 and norm(row[0]['text']).startswith(('ten ban/don vi', 'ten don vi')):
                unit_name = row[1]['text'].strip()
    for kind, obj in raw['order']:
        if kind == 'paragraph':
            text = obj['text'].strip()
            if obj['kind'] == 'watermark': continue
            m = re.match(r'^(SK\s*\d+(?:[.]\d+)?)\s*[—–-]\s*(.+)', text, re.I)
            if m:
                current = new_candidate('master', m[1], m[2])
                candidates.append(current)
                add(current, obj, 'Tên sáng kiến', '6')
            continue
        rows = obj['rows']
        flat = norm(' '.join(b['text'] for r in rows[:3] for b in r))
        if 'ma sang kien/nhiem vu' in flat and 'ten sang kien/nhiem vu' in flat:
            current = new_candidate('detail')
            candidates.append(current)
        # Horizontal Master has a semantic header, not a fixed column count.
        header = next((r for r in rows if len(r) >= 7 and norm(r[0]['text']) in ('ma', 'ma sk', 'ma sang kien')), None)
        if header:
            labels = [b['text'] for b in header]
            for row in rows[rows.index(header) + 1:]:
                if len(row) != len(labels): continue
                values = [b['text'] for b in row]
                name_i = next((i for i, l in enumerate(labels) if norm(l).startswith('sang kien/nhiem vu')), None)
                if name_i is None: continue
                name_raw = values[name_i]
                match = re.search(r'(?:Tên SK|Tên sáng kiến/nhiệm vụ|Tên sáng kiến)\s*:\s*(.*?)(?=\n\s*[-–]?\s*Phạm vi|$)', name_raw, re.S | re.I)
                name = match[1].strip() if match else re.split(r'\n\s*[-–]?\s*Phạm vi\s*:', name_raw, maxsplit=1, flags=re.I)[0].strip()
                if blank(name): continue
                current = new_candidate('master', values[0].strip(), name)
                candidates.append(current)
                for b, label in zip(row, labels):
                    add(current, b, label, label_target(label), 'owner' if norm(label).startswith('owner') else None)
            current = None
            continue
        if current:
            for row in rows:
                # PDF table detection can add empty spacer columns. Collapse only
                # when exactly two nonempty cells remain and the first is a known label.
                nonempty = [b for b in row if b['text'].strip()]
                if len(row) > 2 and len(nonempty) == 2 and norm(nonempty[0]['text']).startswith(('ma sang kien', 'nguon hinh thanh', 'owner')):
                    row = nonempty
                if len(row) == 2:
                    label, value = row[0]['text'], row[1]['text']
                    n = norm(label)
                    if n.startswith('ma sang kien'): current['code'] = '' if blank(value) else value.strip()
                    if n.startswith('ten sang kien'): current['name'] = '' if blank(value) else value.strip()
                    meta = 'owner' if n.startswith('owner') else 'baseline' if n.startswith('baseline') else 'origin' if n.startswith('nguon hinh thanh') else None
                    add(current, row[1], label, label_target(label), meta)
                elif len(row) >= 4 and re.match(r'^(?:\d+[.\s]*)?(hanh vi|co che|ket qua)$', norm(row[0]['text'])):
                    layer = norm(row[0]['text'])
                    target = '7A' if 'hanh vi' in layer else '7B' if 'co che' in layer else '7C'
                    add(current, row[1], row[0]['text'] + ' — Hiện trạng', '3')
                    add(current, row[2], row[0]['text'] + ' — Thay đổi cần đạt', target)
                    add(current, row[3], row[0]['text'] + ' — KPI/bằng chứng', '8')
                else:
                    for b in row:
                        add(current, b, 'Chi tiết chưa xác nhận mapping')
    candidates = [c for c in candidates if c['name'] and not blank(c['name'])]
    # Link early-result rows only by exact normalized name; uncertain links stay unassigned.
    for t in raw['tables']:
        if not t['rows']: continue
        labels = [b['text'] for b in t['rows'][0]]
        linked = next((i for i, l in enumerate(labels) if norm(l) == 'sang kien lien quan'), None)
        if linked is None: continue
        for row in t['rows'][1:]:
            if linked >= len(row): continue
            matches = [c for c in candidates if norm(c['name']).rstrip('.') == norm(row[linked]['text']).rstrip('.')]
            if len(matches) == 1:
                for b, label in zip(row, labels): add(matches[0], b, label, '8')
    filename_code = business_code(filename or path.name)
    for c in candidates:
        c['unit_name'] = unit_name
        if not c['code']: c['issues'].append('missing_code')
        if filename_code and business_code(c['code']) and filename_code != business_code(c['code']):
            c['issues'].append('filename_code_conflict')
        c['mapping_state'] = 'pending'
    used = {b['id'] for c in candidates for b in c['blocks']}
    nfull = norm(full[:20000])
    category = 'forms' if candidates else 'guidance' if 'huong dan' in nfull else 'functions' if 'chuc nang' in nfull and 'nhiem vu' in nfull else 'correspondence' if 'cong van' in nfull or 'quyet dinh' in nfull else 'evidence'
    return {'candidates': candidates, 'unassigned': [b for b in all_blocks if b['id'] not in used],
            'category': category, 'pages': raw['pages'], 'ocr_pages': raw['ocr_pages'],
            'status': 'needs_ocr' if raw['ocr_pages'] else 'parsed' if candidates else 'needs_classification',
            'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}


def rules(record):
    version = record['versions'][-1]
    blocks = {b['id']: b for b in version['blocks']}
    fields = version['fields']
    comments = []
    def issue(rule, step, problem, why, question, evidence):
        comments.append({'id': str(uuid.uuid4()), 'version': version['number'], 'step': step,
            'problem': problem, 'why': why, 'question': question, 'evidence': evidence,
            'priority': 'medium', 'rule': rule, 'state': 'pending', 'reviewer': None,
            'reviewed_at': None, 'internal': True, 'stale': False})
    for step in ('7A', '7B', '7C'):
        if not fields.get(step) and not version.get('revisions', {}).get(step) and step not in version.get('not_applicable', []):
            ev = [b for b in blocks.values() if norm(STEPS[step].replace(' mới', '')) in norm(b['label'])]
            issue('EMPTY_LAYER_V1', step, 'Chưa mapping được nội dung ' + STEPS[step],
                  'Cần đối chiếu phần chưa mapping và giai đoạn triển khai; chưa kết luận đơn vị bỏ trống.',
                  'Nội dung này đã có ở nguồn nào, hay cần bổ sung/làm rõ tính áp dụng?', ev)
    owner = [b for b in blocks.values() if norm(b['label']).startswith('owner') and not blank(b['text'])]
    if not owner:
        issue('OWNER_V1', '8', 'Chưa xác định Owner từ trường nguồn', 'Không tự chọn người thay đơn vị.',
              'Ai chịu trách nhiệm xuyên suốt sáng kiến?', [])
    elif len({norm(b['text']) for b in owner}) > 1 or any(' / ' in b['text'] for b in owner):
        issue('OWNER_CONFLICT_V1', '8', 'Nguồn nêu nhiều giá trị Owner', 'Cần xác nhận đầu mối xuyên suốt.',
              'Các tên này có vai trò gì và ai chịu trách nhiệm đến cùng?', owner)
    text = norm('\n'.join(b['text'] for b in blocks.values()))
    planned = bool(re.search(r'(xac lap|thiet lap|do|xay dung).{0,60}baseline|baseline.{0,70}(30 ngay|thang dau|se xac|thiet lap)', text))
    baseline = [b for b in blocks.values() if 'baseline' in norm(b['label']) and not blank(b['text'])]
    if planned:
        issue('BASELINE_PLAN_V1', '3', 'Có kế hoạch thiết lập baseline, chưa đồng nghĩa đã đo',
              'Phân biệt kế hoạch đo với số liệu nền đã xác minh.', 'Khi nào có số liệu và ai xác nhận nguồn đo?',
              [b for b in blocks.values() if 'baseline' in norm(b['text'])])
    elif not baseline and 'baseline:' not in text:
        issue('BASELINE_REVIEW_V1', '3', 'Chưa tìm được baseline hoặc kế hoạch trong phần đã đọc',
              'Cần đối chiếu mọi nguồn và phần chưa mapping trước khi kết luận thiếu.',
              'Đơn vị đang dùng số liệu nền nào và lấy từ đâu?', [])
    return comments


def export_docx(record, template, destination):
    """Patch only document.xml. Preserve every other original package member."""
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    w = '{' + ns['w'] + '}'
    version = record['versions'][-1]
    blocks = {b['id']: b for b in version['blocks']}
    with ZipFile(template) as zin:
        xml = etree.fromstring(zin.read('word/document.xml'))
        for t in xml.findall('.//w:t', ns):
            for name in ['VNPT TP. Hồ Chí Minh', 'VNPT TP. HCM', 'VNPT TP.HCM', 'TP. Hồ Chí Minh', 'TP. HCM', 'TP.HCM']:
                t.text = (t.text or '').replace(name, record['unit_name'])
        def text_of(e): return '' if e is None else ''.join(e.xpath('.//w:t/text()', namespaces=ns))
        def fill(cell, text):
            proto = cell.find('w:p', ns)
            ppr = deepcopy(proto.find('w:pPr', ns)) if proto is not None and proto.find('w:pPr', ns) is not None else None
            for child in list(cell):
                if child.tag != w + 'tcPr': cell.remove(child)
            for line in text.split('\n') or ['']:
                p = etree.SubElement(cell, w + 'p')
                if ppr is not None: p.append(deepcopy(ppr))
                r = etree.SubElement(p, w + 'r'); t = etree.SubElement(r, w + 't')
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); t.text = line
        tables = xml.findall('.//w:body/w:tbl', ns)
        for row in tables[0].findall('w:tr', ns):
            cells = row.findall('w:tc', ns)
            label = norm(text_of(cells[0]))
            if label == 'ten don vi': fill(cells[1], record['unit_name'])
            elif label.startswith('ten sang kien'): fill(cells[1], record['name'])
            elif label.startswith('nguoi trinh bay'):
                fill(cells[1], '\n\n'.join(b['text'] for b in blocks.values() if norm(b['label']).startswith('owner')))
            elif label.startswith('nhom /'): fill(cells[1], '')
            elif label.startswith(('mat tran vnpt', 'nguon hinh thanh')):
                source_labels = ('cau phan vnpt rising lien ket chinh', 'nguon va lien ket') if label.startswith('mat tran') else ('nguon hinh thanh', 'nguon va lien ket')
                values = [b['text'] for b in blocks.values() if norm(b['label']).startswith(source_labels) and not blank(b['text'])]
                if values: fill(cells[1], text_of(cells[1]) + '\nTheo nguồn (chưa đánh dấu thay đơn vị):\n' + '\n\n'.join(dict.fromkeys(values)))
        for row in tables[1].findall('w:tr', ns)[1:]:
            cells = row.findall('w:tc', ns)
            match = re.match(r'(7[ABC]|[1-8])[.]', text_of(cells[0]))
            if not match: raise ValueError('Mẫu WS1b không khớp nhãn 10 ô trả lời')
            step = match[1]
            answer = version.get('revisions', {}).get(step)
            if answer is None:
                answer = '\n\n'.join(blocks[i]['text'] for i in version['fields'][step] if i in blocks)
            fill(cells[-1], answer)
        for e in xml.findall('.//w:trHeight', ns):
            e.getparent().remove(e)
        # A score is filled only from an explicit submitted unit assessment.
        for group, index, start in [('gate', 2, 0), ('criteria', 3, 1)]:
            for i, row in enumerate(tables[index].findall('w:tr', ns)[start:], 1):
                chosen = version.get('self_rating', {}).get(group, {}).get(str(i))
                if chosen in ('Đạt', 'Chưa đạt', 'Chưa rõ'):
                    fill(row.findall('w:tc', ns)[-1], '\n'.join(
                        ('[x] ' if label == chosen else '[ ] ') + label for label in ('Đạt', 'Chưa đạt', 'Chưa rõ')))
        body = xml.find('w:body', ns)
        status = etree.Element(w + 'p')
        status_run = etree.SubElement(status, w + 'r')
        etree.SubElement(status_run, w + 't').text = ('PHIẾU ĐÃ XÁC NHẬN' if version.get('confirmed') else 'PHIẾU NHÁP — CẦN ĐỐI CHIẾU NGUỒN')
        body.insert(0, status)
        def paragraph(text, page_break=False):
            p = etree.Element(w + 'p')
            if page_break:
                props = etree.SubElement(p, w + 'pPr'); etree.SubElement(props, w + 'pageBreakBefore')
            r = etree.SubElement(p, w + 'r'); t = etree.SubElement(r, w + 't'); t.text = text
            body.insert(len(body) - 1, p)
        paragraph(('PHIẾU ĐÃ XÁC NHẬN' if version.get('confirmed') else 'PHIẾU NHÁP') +
                  f" — {record['code']} — Phiên bản {version['number']}", True)
        paragraph('PHỤ LỤC ĐỐI CHIẾU NGUỒN NGUYÊN VĂN')
        paragraph('Chi tiết dưới đây được bảo toàn để đối chiếu; không phải nhận xét nội bộ hoặc kết quả đã kiểm chứng.')
        for b in blocks.values():
            if blank(b['text']): continue
            paragraph(b.get('file_name', '') + ' — ' + b['label'] + ' — ' + str(b['locator']))
            for line in b['text'].split('\n'): paragraph(line)
            if 'original_text' in b:
                paragraph('Nguyên văn trước hiệu chỉnh:')
                for line in b['original_text'].split('\n'): paragraph(line)
        with ZipFile(destination, 'w') as zout:
            for info in zin.infolist():
                zout.writestr(info, etree.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True)
                              if info.filename == 'word/document.xml' else zin.read(info.filename))
    return str(destination)
